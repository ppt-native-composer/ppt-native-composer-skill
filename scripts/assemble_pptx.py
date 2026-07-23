#!/usr/bin/env python3
"""Assemble an editable V2 PPTX from visual assets and editable layers."""

from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
import zipfile
from pathlib import Path
from typing import Any

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

SCRIPT_DIR = Path(__file__).resolve().parent
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

from validate_blueprint import validate  # noqa: E402
from design_intent_lifecycle import validate_for_evc  # noqa: E402
from evc_runtime import RuntimeAudit, capability_manifest, render_evc  # noqa: E402
from route_normalization import extract_route  # noqa: E402


SLIDE_W = 13.333333
SLIDE_H = 7.5
READY_STATUSES = {"ready", "registered", "exists", "source_ready", "approved"}
CANONICAL_WEIGHT_TOKENS = {"bold", "semibold", "medium", "regular", "light"}
FORBIDDEN_EFFECT_PATTERNS = (
    re.compile(rb"<a:outerShdw\b[^>]*>.*?</a:outerShdw>", re.S),
    re.compile(rb"<a:outerShdw\b[^>]*/>", re.S),
    re.compile(rb"<a:innerShdw\b[^>]*>.*?</a:innerShdw>", re.S),
    re.compile(rb"<a:innerShdw\b[^>]*/>", re.S),
    re.compile(rb"<a:effectDag\b[^>]*>.*?</a:effectDag>", re.S),
    re.compile(rb"<a:effectDag\b[^>]*/>", re.S),
)


def hex_to_rgb(value: str, fallback: str = "#111111") -> RGBColor:
    value = str(value or fallback).strip()
    if not value.startswith("#"):
        value = f"#{value}"
    if len(value) != 7:
        value = fallback
    try:
        return RGBColor(int(value[1:3], 16), int(value[3:5], 16), int(value[5:7], 16))
    except Exception:
        fallback = fallback.lstrip("#")
        return RGBColor(int(fallback[0:2], 16), int(fallback[2:4], 16), int(fallback[4:6], 16))


def zone(value: Any, fallback: dict[str, float]) -> dict[str, float]:
    if not isinstance(value, dict):
        return dict(fallback)
    result = dict(fallback)
    for key in ("x", "y", "w", "h"):
        try:
            result[key] = float(value[key])
        except Exception:
            pass
    return result


def required_zone(value: Any, label: str) -> dict[str, float]:
    if not isinstance(value, dict):
        raise RuntimeError(f"{label}.zone is required")
    return zone(value, {"x": 0, "y": 0, "w": 0, "h": 0})


def visual_bible_colors(data: dict[str, Any]) -> dict[str, str]:
    tokens = data.get("visual_bible", {}).get("design_tokens", {})
    colors = tokens.get("colors") if isinstance(tokens, dict) else None
    if isinstance(colors, dict):
        return {
            "paper": colors.get("paper") or colors.get("background") or "#F7F4EE",
            "background": colors.get("background") or colors.get("paper") or "#F7F4EE",
            "ink": colors.get("ink") or colors.get("text") or "#111111",
            "muted": colors.get("muted") or "#6E6A63",
            "accent": colors.get("accent") or "#E6422E",
            "secondary_accent": colors.get("secondary_accent") or colors.get("accent") or "#276EF1",
            "line": colors.get("line") or "#CFC7BA",
        }
    if isinstance(colors, list):
        raise RuntimeError("visual_bible.design_tokens.colors must be canonical object; legacy list is migration-only")
    return {"paper": "#F7F4EE", "ink": "#111111", "muted": "#6E6A63", "accent": "#E6422E", "line": "#CFC7BA"}


def fonts(data: dict[str, Any]) -> dict[str, str]:
    plan = data.get("font_plan", {})
    return {
        "cjk": str(plan.get("cjk") or "PingFang SC"),
        "latin": str(plan.get("latin") or "Aptos"),
    }


def type_scale(data: dict[str, Any]) -> dict[str, dict[str, Any]]:
    visual_font_plan = data.get("visual_bible", {}).get("font_plan", {})
    scale = visual_font_plan.get("type_scale") if isinstance(visual_font_plan, dict) else None
    if not isinstance(scale, dict):
        plan = data.get("font_plan", {})
        scale = plan.get("type_scale") if isinstance(plan, dict) else None
    default = {
        "display": {"cjk_size_pt": 30, "latin_size_pt": 30, "weight": "bold", "line_height": 1.08, "min_contrast": 4.5},
        "title": {"cjk_size_pt": 24, "latin_size_pt": 24, "weight": "bold", "line_height": 1.1, "min_contrast": 4.5},
        "claim": {"cjk_size_pt": 14, "latin_size_pt": 14, "weight": "semibold", "line_height": 1.15, "min_contrast": 4.5},
        "section": {"cjk_size_pt": 11, "latin_size_pt": 11, "weight": "semibold", "line_height": 1.15, "min_contrast": 4.5},
        "body": {"cjk_size_pt": 10, "latin_size_pt": 10, "weight": "regular", "line_height": 1.25, "min_contrast": 4.5},
        "label": {"cjk_size_pt": 9, "latin_size_pt": 9, "weight": "medium", "line_height": 1.12, "min_contrast": 4.5},
        "caption": {"cjk_size_pt": 8, "latin_size_pt": 8, "weight": "regular", "line_height": 1.2, "min_contrast": 4.5},
        "page_number": {"cjk_size_pt": 7, "latin_size_pt": 7, "weight": "regular", "line_height": 1.0, "min_contrast": 4.5},
    }
    if isinstance(scale, dict):
        for key, value in scale.items():
            if isinstance(value, dict) and key in default:
                default[key] = {**default[key], **value}
    return default


def type_scale_key(item: dict[str, Any], role: str) -> str:
    treatment = item.get("typographic_treatment") if isinstance(item.get("typographic_treatment"), dict) else {}
    explicit = str(treatment.get("type_scale") or treatment.get("scale") or "").strip()
    if explicit:
        return explicit
    level = int(item.get("hierarchy_level") or treatment.get("hierarchy_level") or 4)
    lowered = role.lower()
    if "page_number" in lowered:
        return "page_number"
    if "title" in lowered:
        return "title" if level > 0 else "display"
    if "claim" in lowered or "main_argument" in lowered or "judgment" in lowered:
        return "claim"
    if "section" in lowered:
        return "section"
    if "caption" in lowered or "reasoning" in lowered or "note" in lowered:
        return "caption"
    if "label" in lowered or "keyword" in lowered or "argument" in lowered:
        return "label"
    return "body"


def load_slots(project_dir: Path) -> dict[str, dict[str, Any]]:
    path = project_dir / "assets" / "asset_slots.json"
    if not path.exists():
        return {}
    payload = json.loads(path.read_text(encoding="utf-8"))
    slots = {}
    for item in payload.get("slots", []):
        if isinstance(item, dict):
            slot_id = str(item.get("asset_id") or item.get("id") or "")
            if slot_id:
                slots[slot_id] = item
    return slots


def load_editable_visual_composition(slide_data: dict[str, Any], project_dir: Path) -> dict[str, Any]:
    inline = slide_data.get("editable_visual_composition")
    if isinstance(inline, dict):
        return inline
    ref = str(slide_data.get("editable_visual_composition_ref") or "").strip()
    if not ref:
        return {}
    path = Path(ref)
    if not path.is_absolute():
        path = project_dir / path
    if not path.exists():
        raise RuntimeError(f"editable_visual_composition_ref does not exist: {path}")
    payload = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise RuntimeError(f"editable_visual_composition_ref must contain an object: {path}")
    return payload


def image_has_alpha(path: Path) -> bool:
    with Image.open(path) as img:
        return img.mode in {"RGBA", "LA"} or "transparency" in img.info


def resolve_asset_path(asset: dict[str, Any], slots: dict[str, dict[str, Any]], project_dir: Path) -> tuple[Path, dict[str, Any]]:
    asset_id = str(asset.get("asset_id") or asset.get("id") or "")
    slot = slots.get(asset_id)
    if not slot:
        raise RuntimeError(f"Missing asset slot for {asset_id}")
    status = str(slot.get("status") or "").lower()
    if status not in READY_STATUSES:
        raise RuntimeError(f"Asset {asset_id} is not ready: {status or 'missing'}")
    path = Path(str(slot.get("path") or slot.get("filename") or ""))
    if not path.is_absolute():
        path = project_dir / "assets" / path
    if not path.exists():
        raise RuntimeError(f"Asset file missing for {asset_id}: {path}")
    if slot.get("placeholder_sha256"):
        raise RuntimeError(f"Asset {asset_id} is a placeholder")
    if asset.get("transparency_required") or slot.get("transparency_required"):
        if not image_has_alpha(path):
            raise RuntimeError(f"Asset {asset_id} requires transparency but has no alpha channel")
    return path, slot


def fit_box(path: Path, target: dict[str, float], crop: bool = False) -> dict[str, float]:
    if crop:
        return dict(target)
    with Image.open(path) as img:
        iw, ih = img.size
    if iw <= 0 or ih <= 0:
        return dict(target)
    image_ratio = iw / ih
    box_ratio = target["w"] / target["h"]
    if image_ratio > box_ratio:
        w = target["w"]
        h = w / image_ratio
        x = target["x"]
        y = target["y"] + (target["h"] - h) / 2
    else:
        h = target["h"]
        w = h * image_ratio
        x = target["x"] + (target["w"] - w) / 2
        y = target["y"]
    return {"x": x, "y": y, "w": w, "h": h}


def add_textbox(
    slide,
    *,
    object_id: str,
    role: str,
    text: str,
    box: dict[str, float],
    font_face: str,
    size: int,
    color: str,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
) -> dict[str, Any]:
    shape = slide.shapes.add_textbox(Inches(box["x"]), Inches(box["y"]), Inches(box["w"]), Inches(box["h"]))
    shape.name = object_id
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = Inches(0.03)
    tf.margin_right = Inches(0.03)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    for run in p.runs:
        run.font.name = font_face
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = hex_to_rgb(color)
    return {
        "object_id": object_id,
        "shape_name": shape.name,
        "role": role,
        "editable": True,
        "text": text,
        **box,
    }


def evc_font_size(item: dict[str, Any], role: str, data: dict[str, Any]) -> int:
    scale = type_scale(data)
    key = type_scale_key(item, role)
    spec = scale.get(key, scale["body"])
    return int(float(spec.get("cjk_size_pt") or spec.get("latin_size_pt") or 10))


def normalize_weight_token(value: Any, object_id: str = "") -> str:
    raw = str(value or "").strip().lower().replace("_", "-")
    normalized = raw.replace("semi-bold", "semibold")
    if normalized in CANONICAL_WEIGHT_TOKENS:
        return normalized
    if "semi" in normalized:
        print(f"WARNING: {object_id or 'EVC object'} uses legacy natural-language weight '{raw}'; migrate to weight_token plus weight_note", file=sys.stderr)
        return "semibold"
    if "bold" in normalized:
        print(f"WARNING: {object_id or 'EVC object'} uses legacy natural-language weight '{raw}'; migrate to weight_token plus weight_note", file=sys.stderr)
        return "bold"
    if "medium" in normalized:
        print(f"WARNING: {object_id or 'EVC object'} uses legacy natural-language weight '{raw}'; migrate to weight_token plus weight_note", file=sys.stderr)
        return "medium"
    if "light" in normalized:
        print(f"WARNING: {object_id or 'EVC object'} uses legacy natural-language weight '{raw}'; migrate to weight_token plus weight_note", file=sys.stderr)
        return "light"
    if normalized and normalized != "regular":
        print(f"WARNING: {object_id or 'EVC object'} uses unknown weight '{raw}'; falling back to regular", file=sys.stderr)
    return "regular"


def evc_bold(item: dict[str, Any], role: str, data: dict[str, Any]) -> bool:
    scale = type_scale(data)
    key = type_scale_key(item, role)
    treatment = item.get("typographic_treatment") if isinstance(item.get("typographic_treatment"), dict) else {}
    object_id = str(item.get("id") or item.get("source_id") or role)
    weight = normalize_weight_token(treatment.get("weight_token") or treatment.get("weight") or scale.get(key, {}).get("weight") or "regular", object_id)
    return weight in {"bold", "semibold", "medium"}


def evc_align(item: dict[str, Any]):
    treatment = item.get("typographic_treatment") if isinstance(item.get("typographic_treatment"), dict) else {}
    alignment = str(treatment.get("alignment") or "").lower()
    if "right" in alignment:
        return PP_ALIGN.RIGHT
    if "center" in alignment:
        return PP_ALIGN.CENTER
    return PP_ALIGN.LEFT


def place_evc_text(slide, evc: dict[str, Any], data: dict[str, Any]) -> list[dict[str, Any]]:
    fnts = fonts(data)
    colors = visual_bible_colors(data)
    slide_no = int(evc.get("slide_number") or 0)
    entries: list[dict[str, Any]] = []
    for idx, item in enumerate(evc.get("editable_objects", []) or [], start=1):
        if not isinstance(item, dict):
            continue
        object_id = str(item.get("id") or f"s{slide_no:02d}_evc_{idx}")
        role = str(item.get("role") or "editable_text")
        text = str(item.get("text") or "")
        box = required_zone(item.get("zone"), f"editable_visual_composition.editable_objects[{idx}]")
        entries.append(
            add_textbox(
                slide,
                object_id=object_id,
                role=role,
                text=text,
                box=box,
                font_face=fnts["cjk"],
                size=evc_font_size(item, role, data),
                color=colors["ink"] if role != "page_number" else colors["muted"],
                bold=evc_bold(item, role, data),
                align=evc_align(item),
            )
        )
    return entries


def flatten_editable_items(editable_layer: dict[str, Any]) -> list[tuple[str, dict[str, Any]]]:
    items: list[tuple[str, dict[str, Any]]] = []
    for role in ("body", "labels", "data", "table_text", "notes"):
        value = editable_layer.get(role)
        if isinstance(value, list):
            for item in value:
                if isinstance(item, dict):
                    items.append((role, item))
    return items


def place_editable_text(slide, slide_data: dict[str, Any], data: dict[str, Any], text_zone: dict[str, float]) -> list[dict[str, Any]]:
    fnts = fonts(data)
    colors = visual_bible_colors(data)
    editable = slide_data.get("editable_layer", {})
    slide_no = int(slide_data.get("slide_number") or 0)
    entries: list[dict[str, Any]] = []

    title_box = zone(editable.get("title", {}).get("zone"), {"x": text_zone["x"], "y": text_zone["y"], "w": text_zone["w"], "h": 0.55})
    title_text = str(editable.get("title", {}).get("text") or slide_data.get("title") or "")
    entries.append(add_textbox(slide, object_id=f"s{slide_no:02d}_title", role="title", text=title_text, box=title_box, font_face=fnts["cjk"], size=24, color=colors["ink"], bold=True))

    page_number = editable.get("page_number", {})
    pn_text = str(page_number.get("text") or f"{slide_no:02d}")
    pn_box = zone(page_number.get("zone"), {"x": 12.2, "y": 0.38, "w": 0.55, "h": 0.2})
    entries.append(add_textbox(slide, object_id=f"s{slide_no:02d}_page_number", role="page_number", text=pn_text, box=pn_box, font_face=fnts["latin"], size=8, color=colors["muted"], bold=True, align=PP_ALIGN.RIGHT))

    body_items = flatten_editable_items(editable)
    y = text_zone["y"] + 0.88
    for idx, (role, item) in enumerate(body_items, start=1):
        text = str(item.get("text") or "").strip()
        if not text:
            continue
        default_h = 0.36 if role in {"labels", "data"} else 0.48
        box = zone(item.get("zone"), {"x": text_zone["x"], "y": y, "w": text_zone["w"], "h": default_h})
        object_id = str(item.get("id") or f"s{slide_no:02d}_{role}_{idx}")
        font_size = 10 if role in {"labels", "data"} else 12
        bold = role in {"labels", "data"}
        entries.append(add_textbox(slide, object_id=object_id, role=str(item.get("role") or role), text=text, box=box, font_face=fnts["cjk"], size=font_size, color=colors["ink"], bold=bold))
        y = box["y"] + box["h"] + 0.14
    return entries


def add_background(slide, colors: dict[str, str]) -> None:
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = hex_to_rgb(colors["paper"])


def _place_cover_picture(slide, path: Path, target: dict[str, float]):
    with Image.open(path) as img:
        image_ratio = img.width / img.height
    box_ratio = target["w"] / target["h"]
    picture = slide.shapes.add_picture(
        str(path), Inches(target["x"]), Inches(target["y"]), Inches(target["w"]), Inches(target["h"])
    )
    if image_ratio > box_ratio:
        crop = max(0.0, (1.0 - box_ratio / image_ratio) / 2.0)
        picture.crop_left = crop
        picture.crop_right = crop
    elif image_ratio < box_ratio:
        crop = max(0.0, (1.0 - image_ratio / box_ratio) / 2.0)
        picture.crop_top = crop
        picture.crop_bottom = crop
    return picture


def place_assets(
    slide,
    slide_data: dict[str, Any],
    project_dir: Path,
    slots: dict[str, dict[str, Any]],
    asset_manifest: list[dict[str, Any]],
    audit: RuntimeAudit,
) -> list[dict[str, Any]]:
    plan = slide_data.get("composition_plan", {})
    asset_zone = zone(plan.get("asset_zone"), {"x": 6.5, "y": 1.0, "w": 6.0, "h": 5.5})
    records: list[dict[str, Any]] = []
    assets = []
    for key in ("generated_layer", "sourced_asset_layer"):
        for item in slide_data.get(key, []) or []:
            if isinstance(item, dict):
                assets.append(item)
    layer_rank = {"background_asset": 2, "full_substrate": 2, "element_asset": 3}
    assets.sort(
        key=lambda item: layer_rank.get(
            str(item.get("layer") or ("background_asset" if "substrate" in str(item.get("asset_type") or "") else "element_asset")),
            3,
        )
    )
    for idx, asset in enumerate(assets, start=1):
        path, slot = resolve_asset_path(asset, slots, project_dir)
        target = zone(asset.get("expected_zone"), asset_zone)
        fit_mode = str(asset.get("fit_mode") or plan.get("asset_fit") or "contain").lower()
        fit_mode = {"crop": "cover", "fit": "contain", "explicit": "explicit_fit_box"}.get(fit_mode, fit_mode)
        if fit_mode not in {"contain", "cover", "explicit_fit_box"}:
            raise RuntimeError(f"Unsupported asset fit mode for {asset.get('asset_id')}: {fit_mode}")
        if fit_mode == "cover":
            fitted = dict(target)
            picture = _place_cover_picture(slide, path, target)
            crop_state = {
                "left": float(picture.crop_left),
                "right": float(picture.crop_right),
                "top": float(picture.crop_top),
                "bottom": float(picture.crop_bottom),
            }
        else:
            fitted = dict(target) if fit_mode == "explicit_fit_box" else fit_box(path, target, crop=False)
            picture = slide.shapes.add_picture(str(path), Inches(fitted["x"]), Inches(fitted["y"]), Inches(fitted["w"]), Inches(fitted["h"]))
            crop_state = {"left": 0.0, "right": 0.0, "top": 0.0, "bottom": 0.0}
        picture.name = str(asset.get("asset_id") or f"asset_{idx}")
        rotation = float(asset.get("rotation") or 0.0)
        picture.rotation = rotation
        unsupported_fields: list[str] = []
        requested_opacity = float(asset.get("opacity", 1.0))
        # python-pptx exposes no reliable image-alpha API. Do not silently
        # pretend a requested transparency was applied to a raster picture.
        if requested_opacity != 1.0:
            unsupported_fields.append("opacity")
        requested_layer = str(asset.get("layer") or ("background_asset" if "substrate" in str(asset.get("asset_type") or "") else "element_asset"))
        asset_manifest.append(
            {
                "slide_number": slide_data.get("slide_number"),
                "asset_id": picture.name,
                "asset_type": asset.get("asset_type"),
                "visual_role": asset.get("visual_role"),
                "filename": str(path),
                "status": "placed",
                "slot_status": slot.get("status"),
                "fit_mode": fit_mode,
                "rotation": rotation,
                "crop_state": crop_state,
                "source_path": str(path),
                **fitted,
            }
        )
        record = {
            "source_object_id": picture.name,
            "pptx_shape_name": picture.name,
            "object_type": "asset",
            "asset_id": picture.name,
            "asset_registry_status": slot.get("status"),
            "source_path": str(path),
            "requested_layer": requested_layer,
            "applied_layer": requested_layer,
            "requested_geometry": target,
            "applied_geometry": fitted,
            "requested_fit_mode": fit_mode,
            "applied_fit_mode": fit_mode,
            "rotation": rotation,
            "crop_state": crop_state,
            "warnings": ["raster asset opacity is explicitly unsupported"] if unsupported_fields else [],
            "unsupported_fields": unsupported_fields,
            "editable": True,
        }
        audit.add(record)
        records.append(record)
    return records


def load_design_intent(slide_data: dict[str, Any], project_dir: Path) -> dict[str, Any]:
    inline = slide_data.get("design_intent")
    if isinstance(inline, dict):
        return inline
    ref = str(slide_data.get("design_intent_ref") or "").strip()
    if not ref:
        return {}
    path = Path(ref)
    if not path.is_absolute():
        path = project_dir / path
    if not path.exists():
        raise RuntimeError(f"design_intent_ref does not exist: {path}")
    payload = json.loads(path.read_text(encoding="utf-8"))
    return payload if isinstance(payload, dict) else {}


def prepare_evc_runtime_input(evc: dict[str, Any], data: dict[str, Any]) -> tuple[dict[str, Any], list[str]]:
    migrated = json.loads(json.dumps(evc))
    warnings: list[str] = []
    colors = visual_bible_colors(data)
    for item in migrated.get("editable_objects", []) or []:
        if not isinstance(item, dict):
            continue
        treatment = item.setdefault("typographic_treatment", {})
        role = str(item.get("role") or "editable_text")
        if "cjk_size_pt" not in treatment:
            treatment["cjk_size_pt"] = evc_font_size(item, role, data)
            warnings.append(f"{item.get('id')}: migrated type-scale size into cjk_size_pt")
        treatment.setdefault("latin_size_pt", treatment["cjk_size_pt"])
        if "weight_token" not in treatment:
            treatment["weight_token"] = normalize_weight_token(treatment.get("weight") or "regular", str(item.get("id") or role))
            warnings.append(f"{item.get('id')}: migrated legacy weight into weight_token")
        treatment.setdefault("color", colors["muted"] if role == "page_number" else colors["ink"])
        treatment.setdefault("horizontal_alignment", treatment.get("alignment") or "left")
        treatment.setdefault("vertical_alignment", "top")
        treatment.setdefault("line_break_policy", "preserve")
        treatment.setdefault("bullet", False)
        treatment.setdefault("overflow_policy", {"mode": "warn", "minimum_size_pt": treatment["cjk_size_pt"]})
    return migrated, warnings


def apply_native_graphics(slide, slide_data: dict[str, Any], data: dict[str, Any]) -> list[dict[str, Any]]:
    colors = visual_bible_colors(data)
    entries = []
    slide_no = int(slide_data.get("slide_number") or 0)
    plan = slide_data.get("composition_plan", {})
    text_zone = zone(plan.get("text_safe_zone"), {"x": 0.7, "y": 0.85, "w": 5.45, "h": 5.8})
    for idx, item in enumerate(slide_data.get("native_graphic_layer", []) or [], start=1):
        if not isinstance(item, dict):
            continue
        object_type = str(item.get("object_type") or "line").lower()
        if not str(item.get("reason_required") or "").strip():
            raise RuntimeError(f"native_graphic_layer item {idx} missing reason_required")
        if any(token in object_type for token in ("shadow", "card", "container")):
            raise RuntimeError(f"Forbidden native graphic object_type: {object_type}")
        if "rect" in object_type and "round" not in object_type:
            box = zone(item.get("zone"), {"x": text_zone["x"], "y": text_zone["y"] + text_zone["h"] - 0.1, "w": text_zone["w"], "h": 0.03})
            shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(box["x"]), Inches(box["y"]), Inches(box["w"]), Inches(box["h"]))
            shape.fill.solid()
            shape.fill.fore_color.rgb = hex_to_rgb(colors["accent"])
            shape.line.fill.background()
        else:
            box = zone(item.get("zone"), {"x": text_zone["x"], "y": text_zone["y"] + 0.72, "w": text_zone["w"], "h": 0.0})
            shape = slide.shapes.add_connector(1, Inches(box["x"]), Inches(box["y"]), Inches(box["x"] + box["w"]), Inches(box["y"]))
            shape.line.color.rgb = hex_to_rgb(colors["line"])
            shape.line.width = Pt(0.8)
        shape.name = str(item.get("id") or f"s{slide_no:02d}_native_support_{idx}")
        entries.append({"slide_number": slide_no, "object_id": shape.name, "shape_name": shape.name, "role": "native_graphic_support", "editable": False, "text": "", **box})
    return entries


def apply_evc_native_support(slide, evc: dict[str, Any], data: dict[str, Any]) -> list[dict[str, Any]]:
    colors = visual_bible_colors(data)
    entries = []
    slide_no = int(evc.get("slide_number") or 0)
    for idx, item in enumerate(evc.get("native_support", []) or [], start=1):
        if not isinstance(item, dict):
            continue
        support_type = str(item.get("type") or "line").lower()
        if support_type not in {"line", "tick", "dot", "rule"}:
            raise RuntimeError(f"Unsupported EVC native_support type: {support_type}")
        structural_text = " ".join(str(item.get(key) or "") for key in ("type", "shape_type", "preset", "object_type")).lower()
        if any(token in structural_text for token in ("shadow", "card", "container", "roundrect", "rounded container", "icon cluster")):
            raise RuntimeError(f"Forbidden EVC native_support content: {item.get('id')}")
        box = required_zone(item.get("zone"), f"editable_visual_composition.native_support[{idx}]")
        name = str(item.get("id") or f"s{slide_no:02d}_evc_native_{idx}")
        if support_type == "dot":
            size = 0.035
            shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(box["x"]), Inches(box["y"]), Inches(size), Inches(size))
            shape.fill.solid()
            shape.fill.fore_color.rgb = hex_to_rgb(colors["line"])
            shape.line.fill.background()
            entry_box = {"x": box["x"], "y": box["y"], "w": size, "h": size}
        else:
            width = box["w"] if support_type in {"line", "rule"} else 0.16
            shape = slide.shapes.add_connector(1, Inches(box["x"]), Inches(box["y"]), Inches(box["x"] + width), Inches(box["y"]))
            shape.line.color.rgb = hex_to_rgb(colors["line"])
            shape.line.width = Pt(0.45 if support_type in {"tick", "dot"} else 0.55)
            entry_box = {"x": box["x"], "y": box["y"], "w": width, "h": 0.0}
        shape.name = name
        entries.append({"slide_number": slide_no, "object_id": name, "shape_name": name, "role": "evc_native_support", "editable": False, "text": "", **entry_box})
    return entries


def validate_for_assembly(data: dict[str, Any], blueprint_path: Path, project_dir: Path) -> None:
    errors, warnings = validate(data, mode="assembly", project_dir=project_dir, pptx_path=None)
    if errors:
        joined = "\n".join(f"- {item}" for item in errors)
        raise RuntimeError(f"Assembly validation failed for {blueprint_path}:\n{joined}")
    if not data.get("final_script_source", {}).get("confirmed"):
        raise RuntimeError("final_script_source.confirmed must be true")
    if not isinstance(data.get("visual_bible"), dict) or not data.get("visual_bible"):
        raise RuntimeError("visual_bible is required")
    if warnings:
        print("Assembly validation warnings:", file=sys.stderr)
        for item in warnings:
            print(f"  - {item}", file=sys.stderr)


def strip_forbidden_ooxml_effects(pptx_path: Path) -> None:
    tmp = pptx_path.with_suffix(".nosha.pptx")
    with zipfile.ZipFile(pptx_path, "r") as src, zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            payload = src.read(item.filename)
            if item.filename.startswith("ppt/") and item.filename.endswith(".xml"):
                for pattern in FORBIDDEN_EFFECT_PATTERNS:
                    payload = pattern.sub(b"", payload)
            dst.writestr(item, payload)
    shutil.move(str(tmp), str(pptx_path))


def _resolve_template_layout(prs, layout_index: int | None, layout_name: str | None):
    if layout_name:
        for layout in prs.slide_layouts:
            if layout.name == layout_name:
                return layout
        raise RuntimeError(f"Template layout not found: {layout_name}")
    if layout_index is not None:
        if layout_index < 0 or layout_index >= len(prs.slide_layouts):
            raise RuntimeError(f"Template layout index out of range: {layout_index}")
        return prs.slide_layouts[layout_index]
    # Default blank layout preserves the source master without injecting
    # unplanned placeholders. Placeholder mapping remains an explicit future
    # capability rather than an implicit text fallback.
    return prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0]


def _assemble_impl(
    blueprint_path: Path,
    output_path: Path,
    project_dir: Path,
    template_path: Path | None = None,
    template_layout_index: int | None = None,
    template_layout_name: str | None = None,
) -> dict[str, Any]:
    from assembly_pipeline import assemble_pipeline

    return assemble_pipeline(
        blueprint_path,
        output_path,
        project_dir,
        template_path=template_path,
        template_layout_index=template_layout_index,
        template_layout_name=template_layout_name,
    )


def assemble(
    blueprint_path: Path,
    output_path: Path,
    project_dir: Path,
    template_path: Path | None = None,
    template_layout_index: int | None = None,
    template_layout_name: str | None = None,
) -> dict[str, Any]:
    """Backward-compatible assembly API delegated through assembly_pipeline."""
    from assembly_pipeline import assemble_pipeline

    return assemble_pipeline(
        blueprint_path,
        output_path,
        project_dir,
        template_path=template_path,
        template_layout_index=template_layout_index,
        template_layout_name=template_layout_name,
    )


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("blueprint", help="Path to blueprint JSON.")
    parser.add_argument("-o", "--output", help="Output PPTX path.")
    parser.add_argument("--project-dir", help="Project folder. Defaults to blueprint parent.")
    parser.add_argument("--template", help="Optional supplied PPTX template to preserve masters/layouts.")
    parser.add_argument("--template-layout-index", type=int, help="Template layout index used for assembled slides.")
    parser.add_argument("--template-layout-name", help="Template layout name used for assembled slides.")
    parser.add_argument("--json", action="store_true", help="Print JSON result.")
    args = parser.parse_args()

    blueprint_path = Path(args.blueprint).resolve()
    project_dir = Path(args.project_dir).resolve() if args.project_dir else blueprint_path.parent
    project_dir.mkdir(parents=True, exist_ok=True)
    output_path = Path(args.output).resolve() if args.output else project_dir / "pptx" / f"{blueprint_path.stem}.pptx"

    try:
        result = assemble(
            blueprint_path,
            output_path,
            project_dir,
            Path(args.template).resolve() if args.template else None,
            args.template_layout_index,
            args.template_layout_name,
        )
    except Exception as exc:  # noqa: BLE001
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1

    if args.json:
        print(json.dumps(result, ensure_ascii=False, indent=2))
    else:
        print(f"PPTX: {result['output']}")
        print(f"Slides: {result['slides']}")
        print(f"Asset manifest: {result['asset_manifest']}")
        print(f"Editability map: {result['editability_map']}")
        print(f"Speaker notes: {result['speaker_notes']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
