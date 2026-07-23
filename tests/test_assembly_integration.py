from __future__ import annotations

import copy
import json
import sys
import zipfile
from pathlib import Path

from PIL import Image
from pptx import Presentation

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS))

from assemble_pptx import assemble  # noqa: E402
from test_contract_hardening import base_blueprint  # noqa: E402


def runtime_blueprint() -> dict:
    data = copy.deepcopy(base_blueprint())
    slide = data["slides"][0]
    slide["design_intent"].update(
        {
            "lifecycle_status": "human_approved",
            "route_binding": {"selected_route": "element_asset_hybrid"},
            "approved": True,
        }
    )
    evc = slide["editable_visual_composition"]
    evc["composition_version"] = "v2"
    evc["runtime_version"] = "evc_runtime_v2"
    for item in evc["editable_objects"]:
        treatment = item["typographic_treatment"]
        treatment.update(
            {
                "cjk_font_family": "PingFang SC",
                "latin_font_family": "Aptos",
                "horizontal_alignment": "left",
                "vertical_alignment": "top",
                "line_spacing": 1.1,
                "paragraph_spacing_before_pt": 0,
                "paragraph_spacing_after_pt": 0,
                "margins": {"left": 0.03, "right": 0.03, "top": 0.03, "bottom": 0.03},
                "line_break_policy": "preserve",
                "bullet": False,
                "overflow_policy": {"mode": "warn", "minimum_size_pt": 8},
            }
        )
    slide["generated_layer"][0].update(
        {"expected_zone": {"x": 6, "y": 2, "w": 2, "h": 2}, "fit_mode": "contain"}
    )
    return data


def prepare_project(tmp_path: Path) -> tuple[Path, Path]:
    project = tmp_path / "project"
    assets = project / "assets"
    assets.mkdir(parents=True)
    image_path = assets / "fixture.png"
    Image.new("RGBA", (16, 12), (20, 100, 120, 180)).save(image_path)
    (assets / "asset_slots.json").write_text(
        json.dumps(
            {
                "slots": [
                    {
                        "asset_id": "s01_focus",
                        "path": "fixture.png",
                        "status": "ready",
                        "transparency_required": False,
                    }
                ]
            }
        ),
        encoding="utf-8",
    )
    blueprint = project / "blueprint.json"
    blueprint.write_text(json.dumps(runtime_blueprint(), ensure_ascii=False), encoding="utf-8")
    return project, blueprint


def test_runtime_assembly_writes_editable_pptx_and_v2_reports(tmp_path: Path) -> None:
    project, blueprint = prepare_project(tmp_path)
    output = project / "runtime.pptx"
    result = assemble(blueprint, output, project)
    assert output.exists()
    assert Path(result["editability_map_v2"]).exists()
    runtime = json.loads((project / "composition_runtime_report.json").read_text(encoding="utf-8"))
    assert runtime["slides"][0]["silent_ignore_count"] == 0
    editability = json.loads((project / "editability_map_v2.json").read_text(encoding="utf-8"))
    text_objects = [item for item in editability["objects"] if item.get("object_type") == "editable_text"]
    assert len(text_objects) == 4
    assert all(item["editable"] for item in text_objects)
    with zipfile.ZipFile(output) as package:
        slide_xml = package.read("ppt/slides/slide1.xml").decode("utf-8")
        assert "策略判断页" in slide_xml
        assert "outerShdw" not in slide_xml


def test_asset_cover_fit_is_reported_without_silent_change(tmp_path: Path) -> None:
    project, blueprint = prepare_project(tmp_path)
    data = json.loads(blueprint.read_text(encoding="utf-8"))
    data["slides"][0]["generated_layer"][0]["fit_mode"] = "cover"
    blueprint.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    assemble(blueprint, project / "cover.pptx", project)
    report = json.loads((project / "composition_runtime_report.json").read_text(encoding="utf-8"))
    asset = next(item for item in report["slides"][0]["objects"] if item["object_type"] == "asset")
    assert asset["requested_fit_mode"] == asset["applied_fit_mode"] == "cover"


def test_requested_raster_asset_opacity_is_explicitly_reported(tmp_path: Path) -> None:
    project, blueprint = prepare_project(tmp_path)
    data = json.loads(blueprint.read_text(encoding="utf-8"))
    data["slides"][0]["generated_layer"][0]["opacity"] = 0.35
    blueprint.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    assemble(blueprint, project / "asset_opacity.pptx", project)
    report = json.loads((project / "composition_runtime_report.json").read_text(encoding="utf-8"))
    asset = next(item for item in report["slides"][0]["objects"] if item["object_type"] == "asset")
    assert asset["unsupported_fields"] == ["opacity"]
    assert "explicitly unsupported" in asset["warnings"][0]


def test_supplied_template_master_and_layout_are_preserved_mechanically(tmp_path: Path) -> None:
    project, blueprint = prepare_project(tmp_path)
    template = tmp_path / "supplied_template.pptx"
    template_prs = Presentation()
    template_prs.slides.add_slide(template_prs.slide_layouts[0])
    template_prs.save(template)
    original = Presentation(template)
    output = project / "template_runtime.pptx"
    result = assemble(blueprint, output, project, template_path=template, template_layout_index=0)
    reopened = Presentation(output)
    assert result["template_provenance"]["selected_layout_name"] == original.slide_layouts[0].name
    assert len(reopened.slide_masters) == len(original.slide_masters)
    assert len(reopened.slide_layouts) == len(original.slide_layouts)
    report = json.loads((project / "composition_runtime_report.json").read_text(encoding="utf-8"))
    assert report["template_provenance"]["placeholder_mapping"].startswith("explicitly_unsupported")
