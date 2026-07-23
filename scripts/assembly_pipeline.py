#!/usr/bin/env python3
"""Substantive native-PPT assembly pipeline.

This module owns assembly orchestration while reusing low-level asset, EVC and
primitive helpers from the backwards-compatible `assemble_pptx` module.
"""

from __future__ import annotations

import copy
import json
import os
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.util import Inches


class AssemblyPipelineError(RuntimeError):
    pass


def _backend():
    import assemble_pptx

    return assemble_pptx


def _read_object(value: Path | dict[str, Any] | None, label: str) -> dict[str, Any] | None:
    if value is None:
        return None
    if isinstance(value, dict):
        return copy.deepcopy(value)
    payload = json.loads(Path(value).read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise AssemblyPipelineError(f"{label} must resolve to a JSON object")
    return payload


def load_and_validate_assembly_inputs(blueprint_path: Path, project_dir: Path) -> tuple[dict[str, Any], dict[str, dict[str, Any]]]:
    backend = _backend()
    data = _read_object(blueprint_path, "blueprint")
    assert data is not None
    backend.validate_for_assembly(data, blueprint_path, project_dir)
    return data, backend.load_slots(project_dir)


def resolve_presentation_source(source_pptx: Path | None, template_pptx: Path | None, template_layout_index: int | None, template_layout_name: str | None) -> tuple[Any, Any, dict[str, Any] | None]:
    if source_pptx and template_pptx:
        raise AssemblyPipelineError("source_pptx and template_pptx are mutually exclusive")
    selected = source_pptx or template_pptx
    if selected:
        if not selected.exists():
            raise AssemblyPipelineError(f"presentation source does not exist: {selected}")
        try:
            presentation = Presentation(selected)
        except Exception as exc:  # noqa: BLE001
            raise AssemblyPipelineError(f"invalid presentation source: {selected}: {exc}") from exc
        layout = _backend()._resolve_template_layout(presentation, template_layout_index, template_layout_name)
        provenance = {
            "source_path": str(selected),
            "source_kind": "source_pptx" if source_pptx else "template_pptx",
            "selected_layout_name": layout.name,
            "selected_layout_index": list(presentation.slide_layouts).index(layout),
            "slide_master_count": len(presentation.slide_masters),
            "slide_layout_count": len(presentation.slide_layouts),
            "placeholder_mapping": "explicitly_unsupported_without_a_declared_mapping_plan",
        }
        return presentation, layout, provenance
    presentation = Presentation()
    presentation.slide_width = Inches(_backend().SLIDE_W)
    presentation.slide_height = Inches(_backend().SLIDE_H)
    return presentation, presentation.slide_layouts[6], None


def resolve_route_input(blueprint: dict[str, Any], route_input: Path | dict[str, Any] | None) -> dict[str, Any]:
    backend = _backend()
    payload = _read_object(route_input, "route") if route_input is not None else blueprint
    assert payload is not None
    route = backend.extract_route(payload, strict=False)
    if not route.get("canonical_route"):
        raise AssemblyPipelineError("route input does not resolve to a canonical route")
    for slide in blueprint.get("slides", []) or []:
        if not isinstance(slide, dict):
            continue
        local = backend.extract_route({"page_production_route": slide.get("page_production_route")}, strict=False)
        if local.get("canonical_route") and local["canonical_route"] != route["canonical_route"]:
            raise AssemblyPipelineError(f"route input ({route['canonical_route']}) conflicts with blueprint route ({local['canonical_route']})")
        slide["page_production_route"] = {"selected_route": route["canonical_route"], "approval": {"approved": True}}
    return route


def validate_design_intent_lifecycle(blueprint: dict[str, Any], design_intent_input: Path | dict[str, Any] | None, canonical_route: str) -> dict[str, Any] | None:
    backend = _backend()
    external = _read_object(design_intent_input, "design_intent") if design_intent_input is not None else None
    if external and "design_intent" in external and isinstance(external["design_intent"], dict):
        external = external["design_intent"]
    # Legacy callers historically supplied a one-slide blueprint as the
    # Design Intent source.  Preserve that API only when the intent is
    # unambiguous; multi-slide inputs are rejected rather than guessed.
    if external and isinstance(external.get("slides"), list):
        slides = external["slides"]
        if len(slides) != 1 or not isinstance(slides[0], dict) or not isinstance(slides[0].get("design_intent"), dict):
            raise AssemblyPipelineError("design_intent input must be a Design Intent object or an unambiguous one-slide legacy blueprint")
        external = slides[0]["design_intent"]
    for slide in blueprint.get("slides", []) or []:
        if not isinstance(slide, dict):
            continue
        intent = external or backend.load_design_intent(slide, Path("."))
        if not isinstance(intent, dict) or not intent:
            raise AssemblyPipelineError("design intent input was not consumable")
        errors = backend.validate_for_evc(intent, canonical_route)
        if errors:
            raise AssemblyPipelineError("design intent lifecycle invalid: " + "; ".join(errors))
        slide["design_intent"] = intent
    return external


def initialize_presentation_context(blueprint: dict[str, Any], source_pptx: Path | None, template_pptx: Path | None, template_layout_index: int | None, template_layout_name: str | None) -> dict[str, Any]:
    presentation, layout, provenance = resolve_presentation_source(source_pptx, template_pptx, template_layout_index, template_layout_name)
    backend = _backend()
    return {"presentation": presentation, "layout": layout, "template_provenance": provenance, "colors": backend.visual_bible_colors(blueprint)}


def initialize_runtime_context(blueprint: dict[str, Any]) -> dict[str, Any]:
    backend = _backend()
    return {"fonts": backend.fonts(blueprint), "color_tokens": backend.visual_bible_colors(blueprint), "default_color": backend.visual_bible_colors(blueprint)["ink"], "default_size_pt": 10}


def apply_assets_and_layers(slide, slide_data: dict[str, Any], project_dir: Path, slots: dict[str, dict[str, Any]], asset_manifest: list[dict[str, Any]], audit) -> list[dict[str, Any]]:
    return _backend().place_assets(slide, slide_data, project_dir, slots, asset_manifest, audit)


def dispatch_evc_objects(slide, slide_data: dict[str, Any], blueprint: dict[str, Any], project_dir: Path, runtime_context: dict[str, Any], audit) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    backend = _backend()
    plan = slide_data.get("composition_plan", {})
    text_zone = backend.zone(plan.get("text_safe_zone"), {"x": 0.7, "y": 0.85, "w": 5.45, "h": 5.8})
    evc = backend.load_editable_visual_composition(slide_data, project_dir)
    if not evc:
        return backend.place_editable_text(slide, slide_data, blueprint, text_zone), backend.apply_native_graphics(slide, slide_data, blueprint)
    if evc.get("approval", {}).get("approved") is not True:
        raise AssemblyPipelineError(f"Slide {slide_data.get('slide_number')}: editable_visual_composition.approval.approved must be true")
    runtime_evc, migration_warnings = backend.prepare_evc_runtime_input(evc, blueprint)
    audit.warnings.extend(migration_warnings)
    relation_context = dict(runtime_context)
    relation_context["visual_targets"] = runtime_evc.get("visual_targets", [])
    records = backend.render_evc(slide, runtime_evc, relation_context, audit)
    text_entries: list[dict[str, Any]] = []
    native_entries: list[dict[str, Any]] = []
    for record in records:
        entry = {"object_id": record.get("source_object_id"), "shape_name": record.get("pptx_shape_name"), "role": record.get("role") or record.get("object_type"), "editable": record.get("editable", True), "text": record.get("text", ""), **(record.get("applied_geometry") or {})}
        (text_entries if record.get("object_type") == "editable_text" else native_entries).append(entry)
    return text_entries, native_entries


def finalize_runtime_evidence(project_dir: Path, asset_manifest: list[dict[str, Any]], editability: list[dict[str, Any]], editability_v2: list[dict[str, Any]], runtime_reports: list[dict[str, Any]], notes: list[str], template_provenance: dict[str, Any] | None) -> dict[str, str]:
    backend = _backend()
    def write_atomic(path: Path, payload: str) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        temporary = path.with_name(path.name + ".tmp")
        try:
            temporary.write_text(payload, encoding="utf-8")
            os.replace(temporary, path)
        except Exception:
            temporary.unlink(missing_ok=True)
            raise

    files = {
        "asset_manifest": project_dir / "asset_manifest.json", "editability_map": project_dir / "editability_map.json", "editability_map_v2": project_dir / "editability_map_v2.json",
        "composition_runtime_report": project_dir / "composition_runtime_report.json", "runtime_capability_manifest": project_dir / "runtime_capability_manifest.json", "speaker_notes": project_dir / "speaker_notes.md",
    }
    write_atomic(files["asset_manifest"], json.dumps(asset_manifest, ensure_ascii=False, indent=2))
    write_atomic(files["editability_map"], json.dumps({"objects": editability}, ensure_ascii=False, indent=2))
    write_atomic(files["editability_map_v2"], json.dumps({"runtime_version": "evc_runtime_v2", "objects": editability_v2}, ensure_ascii=False, indent=2))
    write_atomic(files["composition_runtime_report"], json.dumps({"runtime_version": "evc_runtime_v2_phase5r", "template_provenance": template_provenance, "slides": runtime_reports}, ensure_ascii=False, indent=2))
    write_atomic(files["runtime_capability_manifest"], json.dumps(backend.capability_manifest(), ensure_ascii=False, indent=2))
    write_atomic(files["speaker_notes"], "\n".join(notes))
    return {key: str(path) for key, path in files.items()}


def save_presentation_atomic(presentation, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = output_path.with_name(output_path.name + ".tmp")
    if temporary.exists():
        temporary.unlink()
    try:
        presentation.save(temporary)
        _backend().strip_forbidden_ooxml_effects(temporary)
        os.replace(temporary, output_path)
    except Exception:
        temporary.unlink(missing_ok=True)
        raise


def validate_saved_presentation(output_path: Path) -> dict[str, int]:
    try:
        presentation = Presentation(output_path)
    except Exception as exc:  # noqa: BLE001
        raise AssemblyPipelineError(f"saved presentation cannot be reopened: {exc}") from exc
    return {"slide_count": len(presentation.slides), "shape_count": sum(len(slide.shapes) for slide in presentation.slides)}


def return_assembly_result(output_path: Path, evidence: dict[str, str], slide_count: int, provenance: dict[str, Any] | None, route: dict[str, Any], saved: dict[str, int]) -> dict[str, Any]:
    return {"output": str(output_path), "slides": slide_count, **evidence, "template_provenance": provenance, "route_provenance": route, "saved_presentation": saved}


def assemble_pipeline(blueprint_path: Path, output_path: Path, project_dir: Path, *, template_path: Path | None = None, template_layout_index: int | None = None, template_layout_name: str | None = None, source_pptx: Path | None = None, route_input: Path | dict[str, Any] | None = None, design_intent_input: Path | dict[str, Any] | None = None) -> dict[str, Any]:
    backend = _backend()
    blueprint, slots = load_and_validate_assembly_inputs(blueprint_path, project_dir)
    route = resolve_route_input(blueprint, route_input)
    validate_design_intent_lifecycle(blueprint, design_intent_input, route["canonical_route"])
    presentation_context = initialize_presentation_context(blueprint, source_pptx, template_path, template_layout_index, template_layout_name)
    runtime_context = initialize_runtime_context(blueprint)
    presentation = presentation_context["presentation"]
    asset_manifest: list[dict[str, Any]] = []
    editability: list[dict[str, Any]] = []
    editability_v2: list[dict[str, Any]] = []
    runtime_reports: list[dict[str, Any]] = []
    notes: list[str] = []
    for slide_data in blueprint.get("slides", []):
        plan = slide_data.get("composition_plan", {})
        if plan.get("approved") is not True:
            raise AssemblyPipelineError(f"Slide {slide_data.get('slide_number')}: composition_plan.approved must be true")
        slide = presentation.slides.add_slide(presentation_context["layout"])
        if not presentation_context["template_provenance"]:
            backend.add_background(slide, presentation_context["colors"])
        audit = backend.RuntimeAudit(int(slide_data.get("slide_number") or 0), route["canonical_route"])
        assets = apply_assets_and_layers(slide, slide_data, project_dir, slots, asset_manifest, audit)
        editability_v2.extend(assets)
        text_entries, native_entries = dispatch_evc_objects(slide, slide_data, blueprint, project_dir, runtime_context, audit)
        for entry in text_entries:
            entry["slide_number"] = slide_data.get("slide_number")
        editability.extend([{**entry, "slide_number": slide_data.get("slide_number")} for entry in text_entries])
        editability.extend(native_entries)
        editability_v2.extend(audit.objects[len(assets):])
        notes.append(f"## Slide {slide_data.get('slide_number')}: {slide_data.get('title')}\n\n{slide_data.get('speaker_notes', '')}\n")
        runtime_reports.append(audit.to_dict())
    save_presentation_atomic(presentation, output_path)
    saved = validate_saved_presentation(output_path)
    evidence = finalize_runtime_evidence(project_dir, asset_manifest, editability, editability_v2, runtime_reports, notes, presentation_context["template_provenance"])
    return return_assembly_result(output_path, evidence, len(blueprint.get("slides", [])), presentation_context["template_provenance"], route, saved)
