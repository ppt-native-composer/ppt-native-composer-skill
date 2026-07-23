#!/usr/bin/env python3
"""Deterministic editable PowerPoint primitives for the EVC runtime."""

from __future__ import annotations

from typing import Any

from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Inches, Pt

from ooxml_capabilities import OoxmlCapabilityError, set_shape_fill_opacity, set_shape_line_opacity


PRIMITIVE_TYPES = {
    "line",
    "rule",
    "tick",
    "dot",
    "rectangle",
    "ellipse",
    "straight_connector",
    "elbow_connector",
}

FORBIDDEN_APPEARANCE_FIELDS = {"shadow", "glow", "bevel", "soft_edges", "gradient"}


def _rgb(value: str, fallback: str = "#111111") -> RGBColor:
    raw = str(value or fallback).strip().lstrip("#")
    if len(raw) != 6:
        raw = fallback.lstrip("#")
    return RGBColor(int(raw[0:2], 16), int(raw[2:4], 16), int(raw[4:6], 16))


def _geometry(spec: dict[str, Any]) -> dict[str, float]:
    geometry = spec.get("geometry") if isinstance(spec.get("geometry"), dict) else {}
    source = geometry or (spec.get("zone") if isinstance(spec.get("zone"), dict) else {})
    kind = str(source.get("type") or "zone")
    if kind == "axis":
        return {
            "x": float(source.get("x1", 0.0)),
            "y": float(source.get("y1", 0.0)),
            "w": float(source.get("x2", 0.0)) - float(source.get("x1", 0.0)),
            "h": float(source.get("y2", 0.0)) - float(source.get("y1", 0.0)),
        }
    if kind == "point":
        diameter = float(source.get("diameter", source.get("w", 0.04)))
        return {"x": float(source.get("x", 0.0)), "y": float(source.get("y", 0.0)), "w": diameter, "h": float(source.get("h", diameter))}
    return {key: float(source.get(key, 0.0)) for key in ("x", "y", "w", "h")}


def _apply_line(shape, spec: dict[str, Any], unsupported: list[str]) -> None:
    color = spec.get("stroke_color") or spec.get("line_color") or "#777777"
    shape.line.color.rgb = _rgb(str(color))
    shape.line.width = Pt(float(spec.get("stroke_width_pt") or spec.get("line_width_pt") or 0.75))
    dash = str(spec.get("dash_style") or "solid").lower()
    dash_map = {
        "solid": MSO_LINE_DASH_STYLE.SOLID,
        "dash": MSO_LINE_DASH_STYLE.DASH,
        "dot": MSO_LINE_DASH_STYLE.ROUND_DOT,
        "dash_dot": MSO_LINE_DASH_STYLE.DASH_DOT,
    }
    if dash in dash_map:
        shape.line.dash_style = dash_map[dash]
    else:
        unsupported.append("dash_style")
    opacity = float(spec.get("opacity", 1.0))
    if opacity != 1.0:
        try:
            set_shape_line_opacity(shape, opacity)
        except OoxmlCapabilityError:
            unsupported.append("opacity")


def add_primitive(slide, spec: dict[str, Any]) -> tuple[Any, dict[str, Any]]:
    primitive_type = str(spec.get("type") or spec.get("element_type") or "line").lower()
    if primitive_type not in PRIMITIVE_TYPES:
        raise ValueError(f"Unsupported native primitive type: {primitive_type}")
    geometry = _geometry(spec)
    unsupported: list[str] = []
    warnings: list[str] = []
    for field in FORBIDDEN_APPEARANCE_FIELDS:
        if spec.get(field) not in (None, False, "", 0):
            unsupported.append(field)

    if primitive_type in {"line", "rule", "tick", "straight_connector", "elbow_connector"}:
        connector_type = MSO_CONNECTOR.ELBOW if primitive_type == "elbow_connector" else MSO_CONNECTOR.STRAIGHT
        shape = slide.shapes.add_connector(
            connector_type,
            Inches(geometry["x"]),
            Inches(geometry["y"]),
            Inches(geometry["x"] + geometry["w"]),
            Inches(geometry["y"] + geometry["h"]),
        )
        _apply_line(shape, spec, unsupported)
    else:
        shape_type = MSO_SHAPE.OVAL if primitive_type in {"dot", "ellipse"} else MSO_SHAPE.RECTANGLE
        shape = slide.shapes.add_shape(
            shape_type,
            Inches(geometry["x"]),
            Inches(geometry["y"]),
            Inches(geometry["w"]),
            Inches(geometry["h"]),
        )
        fill_color = spec.get("fill_color")
        if fill_color:
            shape.fill.solid()
            shape.fill.fore_color.rgb = _rgb(str(fill_color))
        else:
            shape.fill.background()
        if spec.get("stroke_color") or spec.get("line_color"):
            _apply_line(shape, spec, unsupported)
        else:
            shape.line.fill.background()
        opacity = float(spec.get("opacity", 1.0))
        if opacity != 1.0:
            try:
                set_shape_fill_opacity(shape, opacity)
            except OoxmlCapabilityError:
                unsupported.append("opacity")

    shape.name = str(spec.get("id") or spec.get("element_id") or f"primitive_{len(slide.shapes)}")
    rotation = float(spec.get("rotation") or 0.0)
    shape.rotation = rotation
    applied = {
        "source_object_id": shape.name,
        "pptx_shape_name": shape.name,
        "object_type": primitive_type,
        "requested_layer": spec.get("layer"),
        "applied_layer": spec.get("layer"),
        "applied_z_order": spec.get("z_order"),
        "requested_geometry": spec.get("geometry") or spec.get("zone"),
        "applied_geometry": geometry,
        "requested_appearance": {
            key: spec.get(key)
            for key in ("fill_color", "stroke_color", "line_color", "stroke_width_pt", "dash_style", "opacity", "rotation")
            if key in spec
        },
        "applied_appearance": {
            "rotation": rotation,
            "opacity": float(spec.get("opacity", 1.0)),
        },
        "warnings": warnings,
        "unsupported_fields": sorted(set(unsupported)),
        "editable": True,
    }
    return shape, applied
