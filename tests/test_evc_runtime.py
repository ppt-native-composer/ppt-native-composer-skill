from __future__ import annotations

import sys
from pathlib import Path

import pytest
from pptx import Presentation

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS))

from evc_runtime import RuntimeAudit, render_evc, render_text_object  # noqa: E402


def context() -> dict:
    return {
        "fonts": {"cjk": "PingFang SC", "latin": "Aptos"},
        "color_tokens": {"ink": "#111111"},
        "default_color": "#111111",
        "default_size_pt": 12,
    }


def text_item(mode: str = "warn") -> dict:
    return {
        "id": "title",
        "source_id": "title",
        "role": "title",
        "text": "中文 Title\nSecond line",
        "zone": {"x": 1, "y": 1, "w": 3, "h": 1},
        "layer": "primary_text",
        "typographic_treatment": {
            "cjk_font_family": "PingFang SC",
            "latin_font_family": "Aptos",
            "cjk_size_pt": 24,
            "latin_size_pt": 22,
            "weight_token": "bold",
            "color": "#112233",
            "horizontal_alignment": "left",
            "vertical_alignment": "middle",
            "line_spacing": 1.1,
            "paragraph_spacing_before_pt": 1,
            "paragraph_spacing_after_pt": 2,
            "margins": {"left": 0.04, "right": 0.04, "top": 0.02, "bottom": 0.02},
            "line_break_policy": "preserve",
            "bullet": False,
            "overflow_policy": {"mode": mode, "minimum_size_pt": 12},
        },
    }


def test_requested_typography_is_applied_and_reported() -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    audit = RuntimeAudit(1, "element_asset_hybrid")
    record = render_text_object(slide, text_item(), context(), audit)
    assert record["applied_typography"]["cjk_font_family"] == "PingFang SC"
    assert record["applied_typography"]["weight_token"] == "bold"
    assert record["requested_geometry"] == record["applied_geometry"]
    assert record["unsupported_fields"] == []


def test_overflow_fail_warn_and_shrink_with_floor() -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tiny = text_item("fail")
    tiny["zone"] = {"x": 1, "y": 1, "w": 0.5, "h": 0.1}
    with pytest.raises(ValueError):
        render_text_object(slide, tiny, context(), RuntimeAudit(1))

    warning = text_item("warn")
    warning["zone"] = tiny["zone"]
    warn_record = render_text_object(slide, warning, context(), RuntimeAudit(1))
    assert warn_record["warnings"]

    shrink = text_item("shrink_with_floor")
    shrink["zone"] = {"x": 1, "y": 1, "w": 1.5, "h": 0.45}
    shrink_record = render_text_object(slide, shrink, context(), RuntimeAudit(1))
    assert shrink_record["overflow"]["applied_size"] >= 12
    assert shrink_record["overflow"]["applied_size"] <= 24


def test_paragraph_runs_and_layer_order_are_executed() -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    item = text_item()
    item["paragraphs"] = [
        {"runs": [{"text": "重点", "weight_token": "bold"}, {"text": " detail", "italic": True}]}
    ]
    evc = {
        "native_support": [{"id": "rule", "type": "rule", "zone": {"x": 1, "y": 2, "w": 2, "h": 0}}],
        "editable_objects": [item],
    }
    audit = RuntimeAudit(1)
    records = render_evc(slide, evc, context(), audit)
    assert records[0]["object_type"] == "rule"
    assert records[1]["object_type"] == "editable_text"
    assert slide.shapes[1].text == "重点 detail"


def test_unknown_required_field_is_not_silently_ignored() -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    item = text_item()
    item["typographic_treatment"]["mystery_runtime_field"] = True
    record = render_text_object(slide, item, context(), RuntimeAudit(1))
    assert "typographic_treatment.mystery_runtime_field" in record["unsupported_fields"]


def test_unknown_layer_is_reported_with_deterministic_fallback_rank() -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    item = text_item()
    item["layer"] = "unknown_layer"
    record = render_text_object(slide, item, context(), RuntimeAudit(1))
    assert "layer" in record["unsupported_fields"]
    assert record["applied_z_order"] == 5


def test_textbox_fill_stroke_and_opacity_reporting() -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    item = text_item()
    item["appearance"] = {
        "fill": "#FFFFFF",
        "stroke": "#224466",
        "line_width_pt": 1.25,
        "dash_style": "dash",
        "opacity": 0.5,
    }
    record = render_text_object(slide, item, context(), RuntimeAudit(1))
    assert record["applied_appearance"]["fill"] == "#FFFFFF"
    assert record["applied_appearance"]["stroke"] == "#224466"
    assert record["applied_appearance"]["line_width_pt"] == 1.25
    assert "appearance.opacity" not in record["unsupported_fields"]
    assert record["applied_appearance"]["opacity"] == 0.5


def test_run_level_emphasis_is_executed_without_prebuilt_paragraphs() -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    item = text_item()
    item["text"] = "普通重点文字"
    item["typographic_treatment"]["run_level_emphasis"] = [
        {"text": "重点", "weight_token": "bold", "color": "#AA0000"}
    ]
    render_text_object(slide, item, context(), RuntimeAudit(1))
    runs = slide.shapes[0].text_frame.paragraphs[0].runs
    assert any(run.text == "重点" and run.font.bold for run in runs)
