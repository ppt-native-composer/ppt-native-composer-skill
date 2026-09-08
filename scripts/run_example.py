#!/usr/bin/env python3
"""Run a self-contained technical reconstruction fixture, without client assets."""

from __future__ import annotations

import argparse
import json
import shutil
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation

SCRIPT_DIR = Path(__file__).resolve().parent
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

from assemble_pptx import assemble  # noqa: E402
from inspect_pptx_package import inspect  # noqa: E402
from validate_blueprint import count_ooxml_effects  # noqa: E402

FIXTURE_DIR = SCRIPT_DIR.parent / "fixtures" / "examples" / "approved_artwork"


def run_example(output_dir: Path) -> dict:
    output_dir = output_dir.expanduser().absolute()
    # Refuse even an empty existing folder: no customer assets or previous
    # evidence can be overwritten by this installation check.
    output_dir.mkdir(parents=True, exist_ok=False)
    assets = output_dir / "assets"
    assets.mkdir()
    for name in ("blueprint.json", "copy.json"):
        shutil.copyfile(FIXTURE_DIR / name, output_dir / name)
    shutil.copyfile(FIXTURE_DIR / "asset_slots.json", assets / "asset_slots.json")
    spec = json.loads((FIXTURE_DIR / "swatch.json").read_text(encoding="utf-8"))
    Image.new("RGBA", tuple(spec["size_px"]), tuple(spec["color_rgba"])).save(assets / "swatch.png")

    output = output_dir / "example.pptx"
    result = assemble(output_dir / "blueprint.json", output, output_dir)
    errors, warnings, slide_count = inspect(str(output))
    package = {"result": "FAIL" if errors else "PASS", "errors": errors, "warnings": warnings, "slides": slide_count}
    effects = count_ooxml_effects(output)
    reopened = Presentation(output)
    actual_text = [
        shape.text for slide in reopened.slides for shape in slide.shapes
        if shape.has_text_frame and shape.text
    ]
    expected = json.loads((output_dir / "copy.json").read_text(encoding="utf-8"))
    text_preserved = sorted(actual_text) == sorted(expected.values())
    if errors or any(effects.values()) or not text_preserved:
        raise RuntimeError("Example verification failed; inspect the new output directory.")
    report = {
        "status": "mechanical_pass",
        "fixture_only": True,
        "source": "Bundled original technical swatch and neutral copy; no external assets.",
        "slide_count": len(reopened.slides),
        "editable_text_count": len(actual_text),
        "text_preserved": text_preserved,
        "package_inspection": package,
        "no_shadow_scan": {"status": "pass", "ooxml_effect_counts": effects},
        "python_reopen": "pass",
        "visual_qa": "not_run",
        "client_grade": False,
        "assembly": result,
    }
    (output_dir / "example_result.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    return report


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, required=True, help="New directory; must not already exist.")
    args = parser.parse_args()
    try:
        result = run_example(args.output_dir)
    except (OSError, ValueError, RuntimeError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
