#!/usr/bin/env python3
"""Extract editable text metrics from an existing PPTX.

This helper is read-only: it does not modify the PPTX and does not validate
visual quality. It produces a companion JSON artifact for Layout Intelligence
scorecards.
"""

from __future__ import annotations

import argparse
import json
import math
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.dml import MSO_FILL


EMU_PER_INCH = 914400
PT_PER_INCH = 72
DEFAULT_LINE_HEIGHT = 1.18


def emu_to_inches(value: Any) -> float | None:
    if value is None:
        return None
    try:
        return round(int(value) / EMU_PER_INCH, 4)
    except Exception:
        return None


def pt_value(value: Any) -> float | None:
    if value is None:
        return None
    try:
        return round(float(value.pt), 2)
    except Exception:
        return None


def cjk_latin_punctuation_counts(text: str) -> dict[str, int]:
    cjk = 0
    latin = 0
    punctuation = 0
    for char in text:
        if char.isspace():
            continue
        code = ord(char)
        if 0x4E00 <= code <= 0x9FFF or 0x3400 <= code <= 0x4DBF:
            cjk += 1
        elif char.isascii() and char.isalnum():
            latin += 1
        elif (
            0x3000 <= code <= 0x303F
            or 0xFF00 <= code <= 0xFFEF
            or char in "。，、；：？！“”‘’（）《》—…·,.!?;:()[]{}-"
        ):
            punctuation += 1
        elif char.isascii():
            punctuation += 1
        else:
            cjk += 1
    return {"cjk_char_count": cjk, "latin_char_count": latin, "punctuation_count": punctuation}


def weighted_cjk_count(counts: dict[str, int]) -> float:
    return counts["cjk_char_count"] + counts["punctuation_count"] + counts["latin_char_count"] * 0.5


def normalize_role(shape_name: str) -> str:
    value = shape_name.lower()
    if "title" in value:
        return "title"
    if "claim" in value or "main" in value:
        return "claim"
    if "caption" in value or "body" in value:
        return "caption"
    if "page_number" in value or "page number" in value:
        return "page_number"
    if "proof" in value:
        return "proof_label"
    return "text"


def paragraph_line_count(text: str, paragraph_count: int) -> int:
    explicit_lines = sum(max(1, len(part.splitlines())) for part in text.split("\v"))
    return max(1, paragraph_count, explicit_lines)


def estimate_rendered_lines(text: str, font_size_pt: float | None, bbox_width: float | None) -> tuple[int | None, str]:
    if font_size_pt is None or bbox_width is None or bbox_width <= 0:
        return None, "unavailable"
    explicit_parts = text.splitlines() or [text]
    estimated_lines = 0
    for part in explicit_parts:
        counts = cjk_latin_punctuation_counts(part)
        weighted_chars = weighted_cjk_count(counts)
        avg_char_width_in = (font_size_pt / PT_PER_INCH) * 0.58
        chars_per_line = max(1, int(bbox_width / max(avg_char_width_in, 0.01)))
        estimated_lines += max(1, int((weighted_chars + chars_per_line - 1) // chars_per_line))
    if len(explicit_parts) > 1:
        return estimated_lines, "mixed"
    return estimated_lines, "estimated"


def overflow_status(
    line_count: int | None,
    font_size_pt: float | None,
    bbox_height: float | None,
    line_height: float = DEFAULT_LINE_HEIGHT,
) -> dict[str, Any]:
    if line_count is None or font_size_pt is None or bbox_height is None or bbox_height <= 0:
        return {"status": "not_run", "required_height_inches": None, "available_height_inches": bbox_height}
    required = (font_size_pt / PT_PER_INCH) * line_height * line_count
    ratio = required / bbox_height
    if ratio > 1.08:
        status = "fail"
    elif ratio > 0.92:
        status = "warning"
    else:
        status = "pass"
    return {
        "status": status,
        "required_height_inches": round(required, 4),
        "available_height_inches": bbox_height,
        "height_usage_ratio": round(ratio, 4),
    }


def density_level(density: float | None) -> str:
    if density is None:
        return "not_run"
    if density >= 36:
        return "severe"
    if density >= 24:
        return "high"
    if density >= 14:
        return "medium"
    return "low"


def role_threshold(role: str) -> dict[str, Any]:
    thresholds = {
        "title": {"min": 18, "preferred": (24, 32)},
        "claim": {"min": 12, "preferred": (15, 20)},
        "proof_label": {"min": 9, "preferred": (10.5, 13)},
        "caption": {"min": 7.5, "preferred": (8.5, 10)},
        "page_number": {"min": 7.5, "preferred": (8.5, 10)},
    }
    return thresholds.get(role, {"min": 8, "preferred": (9, 12)})


def preferred_range_status(role: str, font_size_pt: float | None) -> str:
    if font_size_pt is None:
        return "not_run"
    cfg = role_threshold(role)
    if font_size_pt < cfg["min"]:
        return "below_min"
    low, high = cfg["preferred"]
    if font_size_pt < low:
        return "below_preferred"
    if font_size_pt <= high:
        return "within_preferred"
    return "above_preferred"


def font_color(shape: Any) -> dict[str, Any]:
    colors: list[str] = []
    if not getattr(shape, "has_text_frame", False):
        return {"text_color": None, "status": "not_run"}
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            color = run.font.color
            try:
                rgb = color.rgb
            except Exception:
                rgb = None
            if rgb is not None:
                colors.append(f"#{str(rgb)}")
    unique = sorted(set(colors))
    if len(unique) == 1:
        return {
            "text_color": unique[0],
            "dominant_text_color": unique[0],
            "has_mixed_text_colors": False,
            "text_color_source": "run",
            "status": "available",
        }
    if len(unique) > 1:
        return {
            "text_color": None,
            "dominant_text_color": unique[0],
            "text_colors": unique,
            "has_mixed_text_colors": True,
            "text_color_source": "run",
            "status": "mixed",
        }
    return {
        "text_color": None,
        "dominant_text_color": None,
        "has_mixed_text_colors": False,
        "text_color_source": "unknown",
        "status": "not_run",
    }


def fill_hex(fill: Any) -> str | None:
    try:
        if fill.type != MSO_FILL.SOLID:
            return None
        rgb = fill.fore_color.rgb
    except Exception:
        return None
    if rgb is None:
        return None
    return f"#{str(rgb)}"


def slide_background_color(slide: Any) -> str | None:
    try:
        return fill_hex(slide.background.fill)
    except Exception:
        return None


def shape_fill_color(shape: Any) -> str | None:
    try:
        return fill_hex(shape.fill)
    except Exception:
        return None


def hex_to_rgb(value: str | None) -> tuple[float, float, float] | None:
    if not value:
        return None
    value = value.strip().lstrip("#")
    if len(value) != 6:
        return None
    try:
        return int(value[0:2], 16) / 255, int(value[2:4], 16) / 255, int(value[4:6], 16) / 255
    except Exception:
        return None


def relative_luminance(rgb: tuple[float, float, float]) -> float:
    def channel(value: float) -> float:
        return value / 12.92 if value <= 0.03928 else ((value + 0.055) / 1.055) ** 2.4

    r, g, b = (channel(item) for item in rgb)
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def contrast_ratio(foreground: str | None, background: str | None) -> float | None:
    fg = hex_to_rgb(foreground)
    bg = hex_to_rgb(background)
    if fg is None or bg is None:
        return None
    lum1 = relative_luminance(fg)
    lum2 = relative_luminance(bg)
    lighter = max(lum1, lum2)
    darker = min(lum1, lum2)
    return round((lighter + 0.05) / (darker + 0.05), 2)


def contrast_threshold_for_role(role: str, font_size_pt: float | None, bold: bool | None) -> float:
    if role == "page_number":
        return 3.0
    if font_size_pt is not None:
        if font_size_pt >= 18 or (bold is True and font_size_pt >= 14):
            return 3.0
    return 4.5


def contrast_status_for(ratio: float | None, threshold: float, warning_margin: float = 0.3) -> str:
    if ratio is None:
        return "not_run"
    if ratio >= threshold:
        return "pass"
    if ratio >= threshold - warning_margin:
        return "warning"
    return "fail"


def bboxes_overlap(a: dict[str, float | None], b: dict[str, Any]) -> bool:
    try:
        ax1 = float(a["x"])
        ay1 = float(a["y"])
        ax2 = ax1 + float(a["w"])
        ay2 = ay1 + float(a["h"])
        bx1 = float(b["x"])
        by1 = float(b["y"])
        bx2 = bx1 + float(b["w"])
        by2 = by1 + float(b["h"])
    except Exception:
        return False
    return ax1 < bx2 and ax2 > bx1 and ay1 < by2 and ay2 > by1


def load_texture_boxes(path: str | None) -> list[dict[str, Any]]:
    if not path:
        return []
    p = Path(path)
    if not p.exists():
        return []
    try:
        payload = json.loads(p.read_text(encoding="utf-8"))
    except Exception:
        return []
    boxes = []
    for asset_id, asset in (payload.get("selected_local_assets") or {}).items():
        if isinstance(asset, dict) and isinstance(asset.get("bbox"), dict):
            bbox = dict(asset["bbox"])
            bbox["asset_id"] = asset_id
            bbox["background_type"] = "image_or_texture"
            boxes.append(bbox)
    return boxes


def background_metrics(shape: Any, slide: Any, bbox: dict[str, float | None], texture_boxes: list[dict[str, Any]]) -> dict[str, Any]:
    overlapping_textures = [item for item in texture_boxes if bboxes_overlap(bbox, item)]
    shape_fill = shape_fill_color(shape)
    slide_bg = slide_background_color(slide)
    if overlapping_textures:
        return {
            "background_type": "image_or_texture",
            "shape_fill_color": shape_fill,
            "slide_background_color": slide_bg,
            "background_color": None,
            "nearest_fill_color": None,
            "overlapping_texture_assets": [item.get("asset_id") for item in overlapping_textures],
            "contrast_status_base": "manual_required",
        }
    if shape_fill:
        return {
            "background_type": "shape_solid_fill",
            "shape_fill_color": shape_fill,
            "slide_background_color": slide_bg,
            "background_color": shape_fill,
            "nearest_fill_color": shape_fill,
            "overlapping_texture_assets": [],
            "contrast_status_base": None,
        }
    if slide_bg:
        return {
            "background_type": "slide_solid_fill",
            "shape_fill_color": shape_fill,
            "slide_background_color": slide_bg,
            "background_color": slide_bg,
            "nearest_fill_color": slide_bg,
            "overlapping_texture_assets": [],
            "contrast_status_base": None,
        }
    return {
        "background_type": "unknown",
        "shape_fill_color": shape_fill,
        "slide_background_color": slide_bg,
        "background_color": None,
        "nearest_fill_color": None,
        "overlapping_texture_assets": [],
        "contrast_status_base": "not_run",
    }


def font_metrics(shape: Any) -> dict[str, Any]:
    sizes: list[float] = []
    bold_values: list[bool] = []
    italic_values: list[bool] = []
    names: list[str] = []
    run_count = 0
    paragraph_count = 0
    if not getattr(shape, "has_text_frame", False):
        return {
            "font_size_pt": None,
            "font_sizes_pt": [],
            "has_mixed_font_sizes": False,
            "font_weight": None,
            "bold": None,
            "italic": None,
            "font_name": None,
            "paragraph_count": 0,
            "run_count": 0,
            "warnings": ["shape has no text frame"],
        }
    for paragraph in shape.text_frame.paragraphs:
        paragraph_count += 1
        for run in paragraph.runs:
            run_count += 1
            size = pt_value(run.font.size)
            if size is not None:
                sizes.append(size)
            if run.font.bold is not None:
                bold_values.append(bool(run.font.bold))
            if run.font.italic is not None:
                italic_values.append(bool(run.font.italic))
            if run.font.name:
                names.append(str(run.font.name))
    unique_sizes = sorted(set(sizes))
    unique_names = sorted(set(names))
    bold = None if not bold_values else any(bold_values)
    italic = None if not italic_values else any(italic_values)
    if bold is True:
        font_weight = "bold"
    elif bold is False:
        font_weight = "regular"
    else:
        font_weight = None
    warnings: list[str] = []
    if not sizes:
        warnings.append("missing font_size_pt")
    if not bold_values:
        warnings.append("missing bold metric")
    return {
        "font_size_pt": max(sizes) if sizes else None,
        "font_sizes_pt": unique_sizes,
        "has_mixed_font_sizes": len(unique_sizes) > 1,
        "font_weight": font_weight,
        "bold": bold,
        "italic": italic,
        "font_name": unique_names[0] if len(unique_names) == 1 else None,
        "font_names": unique_names,
        "paragraph_count": paragraph_count,
        "run_count": run_count,
        "warnings": warnings,
    }


def text_frame_metrics(shape: Any) -> dict[str, Any]:
    if not getattr(shape, "has_text_frame", False):
        return {
            "vertical_alignment": None,
            "margins": {"left": None, "right": None, "top": None, "bottom": None},
        }
    frame = shape.text_frame
    vertical = str(frame.vertical_anchor).split(".")[-1] if frame.vertical_anchor is not None else None
    return {
        "vertical_alignment": vertical,
        "margins": {
            "left": emu_to_inches(frame.margin_left),
            "right": emu_to_inches(frame.margin_right),
            "top": emu_to_inches(frame.margin_top),
            "bottom": emu_to_inches(frame.margin_bottom),
        },
    }


def shape_record(slide_number: int, slide: Any, shape: Any, texture_boxes: list[dict[str, Any]]) -> dict[str, Any] | None:
    if not getattr(shape, "has_text_frame", False):
        return None
    text = shape.text or ""
    if not text.strip():
        return None
    shape_name = str(shape.name)
    bbox = {
        "x": emu_to_inches(shape.left),
        "y": emu_to_inches(shape.top),
        "w": emu_to_inches(shape.width),
        "h": emu_to_inches(shape.height),
    }
    area = None
    if bbox["w"] is not None and bbox["h"] is not None:
        area = round(bbox["w"] * bbox["h"], 4)
    counts = cjk_latin_punctuation_counts(text)
    density = None
    if area and area > 0:
        density = round(weighted_cjk_count(counts) / area, 4)
    fonts = font_metrics(shape)
    frame = text_frame_metrics(shape)
    line_count = paragraph_line_count(text, fonts["paragraph_count"])
    rendered_line_count, line_count_source = estimate_rendered_lines(text, fonts["font_size_pt"], bbox["w"])
    overflow = overflow_status(rendered_line_count, fonts["font_size_pt"], bbox["h"])
    role = normalize_role(shape_name)
    preferred_status = preferred_range_status(role, fonts["font_size_pt"])
    colors = font_color(shape)
    background = background_metrics(shape, slide, bbox, texture_boxes)
    ratio = None
    contrast_status = background.get("contrast_status_base")
    if contrast_status is None:
        ratio = contrast_ratio(colors.get("dominant_text_color") or colors.get("text_color"), background.get("background_color"))
        contrast_status = contrast_status_for(ratio, contrast_threshold_for_role(role, fonts["font_size_pt"], fonts["bold"]))
    overflow_risk = None
    if area and density is not None:
        overflow_risk = density > 36 or line_count > 4
    warnings = list(fonts.pop("warnings"))
    if bbox["w"] is None or bbox["h"] is None:
        warnings.append("missing bbox metric")
    return {
        "slide_number": slide_number,
        "object_id": shape_name,
        "shape_name": shape_name,
        "role": role,
        "editable": True,
        "text": text,
        "bbox": bbox,
        "font_size_pt": fonts["font_size_pt"],
        "font_weight": fonts["font_weight"],
        "bold": fonts["bold"],
        "italic": fonts["italic"],
        "font_name": fonts["font_name"],
        "font_names": fonts["font_names"],
        "line_count_estimated": line_count,
        "estimated_rendered_line_count_v2": rendered_line_count,
        "line_count_source": line_count_source,
        "paragraph_count": fonts["paragraph_count"],
        "run_count": fonts["run_count"],
        "cjk_char_count": counts["cjk_char_count"],
        "latin_char_count": counts["latin_char_count"],
        "punctuation_count": counts["punctuation_count"],
        "weighted_cjk_char_count": weighted_cjk_count(counts),
        "text_box_area": area,
        "cjk_density": density,
        "text_density_level": density_level(density),
        "has_mixed_font_sizes": fonts["has_mixed_font_sizes"],
        "role_minimum_violation": preferred_status == "below_min",
        "role_preferred_range_status": preferred_status,
        "text_color": colors.get("text_color"),
        "dominant_text_color": colors.get("dominant_text_color"),
        "has_mixed_text_colors": colors.get("has_mixed_text_colors"),
        "text_color_source": colors.get("text_color_source"),
        "text_color_status": colors.get("status"),
        "background_type": background.get("background_type"),
        "shape_fill_color": background.get("shape_fill_color"),
        "slide_background_color": background.get("slide_background_color"),
        "background_color": background.get("background_color"),
        "nearest_fill_color": background.get("nearest_fill_color"),
        "overlapping_texture_assets": background.get("overlapping_texture_assets"),
        "contrast_ratio": ratio,
        "contrast_status": contrast_status,
        "bound_element_ids": [],
        "vertical_alignment": frame["vertical_alignment"],
        "margins": frame["margins"],
        "overflow_risk_estimate": overflow_risk,
        "overflow_risk": overflow["status"],
        "overflow_required_height_inches": overflow["required_height_inches"],
        "overflow_available_height_inches": overflow["available_height_inches"],
        "overflow_height_usage_ratio": overflow.get("height_usage_ratio"),
        "warnings": warnings,
    }


def extract_metrics(pptx: Path, assembly_plan: str | None = None) -> dict[str, Any]:
    presentation = Presentation(str(pptx))
    texture_boxes = load_texture_boxes(assembly_plan)
    objects: list[dict[str, Any]] = []
    warnings: list[str] = []
    for slide_index, slide in enumerate(presentation.slides, start=1):
        for shape in slide.shapes:
            record = shape_record(slide_index, slide, shape, texture_boxes)
            if record:
                objects.append(record)
                for warning in record["warnings"]:
                    warnings.append(f"{record['shape_name']}: {warning}")
    return {
        "source_pptx": str(pptx),
        "assembly_plan": assembly_plan,
        "schema": "editability_metrics_v0_1",
        "object_count": len(objects),
        "objects": objects,
        "warnings": warnings,
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--pptx", required=True)
    parser.add_argument("--assembly-plan")
    parser.add_argument("--output", required=True)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    pptx = Path(args.pptx).resolve()
    if not pptx.exists():
        raise SystemExit(f"PPTX does not exist: {pptx}")
    output = Path(args.output).resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    payload = extract_metrics(pptx, args.assembly_plan)
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Wrote editability metrics: {output}")
    print(f"Objects: {payload['object_count']}")
    print(f"Warnings: {len(payload['warnings'])}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
