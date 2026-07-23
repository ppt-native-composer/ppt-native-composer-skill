from __future__ import annotations

import json
import sys
from pathlib import Path

import pytest
from pptx import Presentation

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS))

from deck_orchestrator import DeckError, DeckRunner, merge_single_page_pptx, selected_page_ids, topological_levels, validate_manifest  # noqa: E402
from runtime_test_support import create_runtime_case  # noqa: E402
import run_deck  # noqa: E402


def _manifest(tmp_path: Path, *, policy: str = "strict", page_count: int = 3) -> Path:
    pages = []
    for number in range(1, page_count + 1):
        project, runtime_manifest = create_runtime_case(tmp_path / f"page{number}", case_id=f"page{number}")
        pages.append({
            "page_id": f"page{number}", "order": number, "section": "test", "page_role": "strategy" if number == 1 else "concept",
            "page_runtime_manifest": str(runtime_manifest), "dependencies": [f"page{number - 1}"] if number == 3 else [],
            "required": number != page_count, "continuity_group": "synthetic", "route": "element_asset_hybrid",
        })
    payload = {
        "manifest_version": "1.0", "deck_id": "deck_test", "deck_title": "Deck test", "source_provenance": {"kind": "synthetic"},
        "page_numbering": {"policy": "manifest_order"}, "failure_policy": policy, "pages": pages,
        "output": {"root": str(tmp_path / "deck_runtime")}, "runtime_options": {"resume": True, "jobs": 2, "preview_dpi": 72},
    }
    path = tmp_path / "deck_manifest.json"
    path.write_text(json.dumps(payload), encoding="utf-8")
    return path


def test_manifest_graph_and_cycle_contract(tmp_path: Path) -> None:
    manifest = json.loads(_manifest(tmp_path).read_text())
    assert validate_manifest(manifest) == []
    assert topological_levels({"a": [], "b": [], "c": ["a", "b"]}) == [["a", "b"], ["c"]]
    assert selected_page_ids({"a": [], "b": ["a"], "c": ["b"]}, "b", None, ["a", "b", "c"]) == {"b", "c"}
    with pytest.raises(DeckError, match="cycle"):
        topological_levels({"a": ["b"], "b": ["a"]})
    manifest["pages"][1]["page_id"] = manifest["pages"][0]["page_id"]
    assert "duplicate page_id" in validate_manifest(manifest)


def test_real_runtime_pages_merge_and_targeted_reuse(tmp_path: Path) -> None:
    manifest = _manifest(tmp_path)
    first = DeckRunner(manifest, skip_preview=True).run()
    assert first["overall_status"] == "success"
    deck = Path(first["deck_pptx"])
    assert len(Presentation(deck).slides) == 3
    first_index = json.loads((Path(first["pages"][0]["page_run_root"]) / "run_summary.json").read_text())
    assert first_index["overall_status"] == "success"
    second = DeckRunner(manifest, page="page2", skip_preview=True).run()
    assert second["overall_status"] == "success"
    assert len(Presentation(second["deck_pptx"]).slides) == 3
    rows = {item["page_id"]: item for item in second["pages"]}
    assert rows["page1"]["cache_status"] == "deck_reuse"
    assert rows["page3"]["cache_status"] != "deck_reuse"


def test_isolate_optional_failure_keeps_auditable_partial_deck(tmp_path: Path) -> None:
    manifest = _manifest(tmp_path, policy="isolate_and_continue")
    data = json.loads(manifest.read_text())
    optional = data["pages"][-1]
    optional["dependencies"] = []
    optional_runtime = Path(optional["page_runtime_manifest"])
    runtime_data = json.loads(optional_runtime.read_text())
    runtime_data["inputs"]["evc"] = str(tmp_path / "missing.json")
    optional_runtime.write_text(json.dumps(runtime_data), encoding="utf-8")
    manifest.write_text(json.dumps(data), encoding="utf-8")
    result = DeckRunner(manifest, skip_preview=True).run()
    assert result["overall_status"] == "partial"
    assert Path(result["deck_pptx"]).exists()
    assert len(Presentation(result["deck_pptx"]).slides) == 2
    assert next(item for item in result["pages"] if item["page_id"] == "page3")["status"] == "failed"


def test_strict_required_failure_does_not_publish_deck(tmp_path: Path) -> None:
    manifest = _manifest(tmp_path, policy="strict")
    data = json.loads(manifest.read_text())
    runtime = Path(data["pages"][0]["page_runtime_manifest"])
    broken = json.loads(runtime.read_text())
    broken["inputs"]["evc"] = str(tmp_path / "missing.json")
    runtime.write_text(json.dumps(broken), encoding="utf-8")
    manifest.write_text(json.dumps(data), encoding="utf-8")
    result = DeckRunner(manifest, skip_preview=True).run()
    assert result["overall_status"] == "failed"
    assert result["deck_pptx"] is None


def test_merge_reopens_and_preserves_slide_shapes(tmp_path: Path) -> None:
    _, first_manifest = create_runtime_case(tmp_path / "one", case_id="one")
    _, second_manifest = create_runtime_case(tmp_path / "two", case_id="two")
    from run_runtime_case import RuntimeCaseRunner
    first, second = RuntimeCaseRunner(first_manifest, skip_preview=True).run(), RuntimeCaseRunner(second_manifest, skip_preview=True).run()
    pages = [Path(next(stage for stage in row["stages"] if stage["stage_id"] == "assemble_pptx")["artifacts"][0]["path"]) for row in (first, second)]
    merged = tmp_path / "merged.pptx"
    mapping = merge_single_page_pptx(pages, merged)
    presentation = Presentation(merged)
    assert len(mapping) == len(presentation.slides) == 2
    assert all(len(slide.shapes) > 0 for slide in presentation.slides)


def test_deck_cli_dry_plan_and_error_exit(tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]) -> None:
    manifest = _manifest(tmp_path)
    monkeypatch.setattr(sys, "argv", ["run_deck.py", "--manifest", str(manifest), "--dry-plan", "--json-summary"])
    assert run_deck.main() == 0
    assert '"dry_plan": true' in capsys.readouterr().out
    monkeypatch.setattr(sys, "argv", ["run_deck.py", "--manifest", str(tmp_path / "missing.json")])
    assert run_deck.main() == 1
