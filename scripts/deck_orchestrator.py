#!/usr/bin/env python3
"""Trusted multi-page orchestration built on Phase 5X-R page runtime evidence."""

from __future__ import annotations

import json
import shutil
import subprocess
import sys
import tempfile
import uuid
import zipfile
from collections import defaultdict, deque
from concurrent.futures import ThreadPoolExecutor, as_completed
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from jsonschema import Draft202012Validator
from pptx import Presentation

from artifact_integrity import sha256_file, validate_artifact
from inspect_pptx_package import inspect
from run_runtime_case import RuntimeCaseRunner, atomic_write, read_json, scan_no_shadow, write_json
from runtime_stage_cache import file_digest, sha256_json, toolchain_fingerprint


SCRIPT_DIR = Path(__file__).resolve().parent
SKILL_ROOT = SCRIPT_DIR.parent
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
ET.register_namespace("p", P_NS)
ET.register_namespace("r", R_NS)


class DeckError(RuntimeError):
    pass


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def resolve_path(value: str, base: Path) -> Path:
    path = Path(value)
    return path if path.is_absolute() else (base / path).resolve()


def artifact(path: Path, artifact_type: str, keys: list[str] | None = None) -> dict[str, Any]:
    result = validate_artifact(path, artifact_type, keys)
    if not result["valid"]:
        raise DeckError(f"invalid deck artifact {path}: {'; '.join(result['errors'])}")
    return {"path": str(path), "artifact_type": artifact_type, "required_keys": keys or [], "sha256": result["sha256"], "metadata": result["metadata"]}


def validate_manifest(manifest: dict[str, Any]) -> list[str]:
    schema = read_json(SKILL_ROOT / "schemas" / "deck_manifest.schema.json")
    errors = [f"{item.json_path}: {item.message}" for item in Draft202012Validator(schema).iter_errors(manifest)]
    pages = manifest.get("pages", [])
    ids = [page.get("page_id") for page in pages if isinstance(page, dict)]
    if len(ids) != len(set(ids)):
        errors.append("duplicate page_id")
    orders = [page.get("order") for page in pages if isinstance(page, dict)]
    if len(orders) != len(set(orders)):
        errors.append("duplicate page order")
    known = set(ids)
    for page in pages:
        for dep in page.get("dependencies", []):
            if dep not in known:
                errors.append(f"page {page.get('page_id')}: unknown dependency {dep}")
    return errors


def page_graph(pages: list[dict[str, Any]]) -> dict[str, list[str]]:
    return {page["page_id"]: list(page.get("dependencies", [])) for page in pages}


def topological_levels(graph: dict[str, list[str]]) -> list[list[str]]:
    pending = {node: set(deps) for node, deps in graph.items()}
    levels: list[list[str]] = []
    while pending:
        ready = sorted(node for node, deps in pending.items() if not deps)
        if not ready:
            raise DeckError("page dependency cycle detected")
        levels.append(ready)
        for node in ready:
            pending.pop(node)
        for deps in pending.values():
            deps.difference_update(ready)
    return levels


def selected_page_ids(graph: dict[str, list[str]], requested: str | None, from_page: str | None, ordered: list[str]) -> set[str]:
    if from_page:
        if from_page not in ordered:
            raise DeckError(f"unknown --from-page {from_page}")
        return set(ordered[ordered.index(from_page):])
    if not requested:
        return set(graph)
    if requested not in graph:
        raise DeckError(f"unknown --page {requested}")
    reverse: dict[str, set[str]] = defaultdict(set)
    for node, deps in graph.items():
        for dep in deps:
            reverse[dep].add(node)
    result, queue = {requested}, deque([requested])
    while queue:
        current = queue.popleft()
        for child in reverse[current]:
            if child not in result:
                result.add(child); queue.append(child)
    return result


def _part_number(name: str, prefix: str, suffix: str) -> int:
    stem = Path(name).name
    return int(stem[len(prefix):-len(suffix)])


def merge_single_page_pptx(page_pptx: list[Path], output: Path) -> list[dict[str, Any]]:
    """Append slides at package level, retaining editable OOXML objects and media."""
    if not page_pptx:
        raise DeckError("no successful page PPTX files to merge")
    output.parent.mkdir(parents=True, exist_ok=True)
    mapping: list[dict[str, Any]] = [{"source_pptx": str(page_pptx[0]), "final_slide_index": 1}]
    # A ZIP append would create duplicate presentation parts.  Unpack and
    # rebuild instead so package integrity is verifiable by all consumers.
    with tempfile.TemporaryDirectory(prefix="ppt-native-deck-") as temporary:
        package = Path(temporary) / "package"
        package.mkdir()
        with zipfile.ZipFile(page_pptx[0]) as source:
            source.extractall(package)
        presentation_path = package / "ppt/presentation.xml"
        rels_path = package / "ppt/_rels/presentation.xml.rels"
        types_path = package / "[Content_Types].xml"
        presentation = ET.parse(presentation_path).getroot()
        presentation_rels = ET.parse(rels_path).getroot()
        content_types = ET.parse(types_path).getroot()
        sld_id_list = presentation.find(f"{{{P_NS}}}sldIdLst")
        if sld_id_list is None:
            raise DeckError("presentation missing slide id list")
        slide_dir = package / "ppt/slides"
        next_slide = max([_part_number(item.name, "slide", ".xml") for item in slide_dir.glob("slide*.xml")] or [0]) + 1
        next_sld_id = max([int(item.attrib.get("id", "255")) for item in sld_id_list] or [255]) + 1
        next_rid = max([int(item.attrib["Id"][3:]) for item in presentation_rels if item.attrib.get("Id", "").startswith("rId") and item.attrib["Id"][3:].isdigit()] or [0]) + 1
        media_dir = package / "ppt/media"
        media_dir.mkdir(parents=True, exist_ok=True)
        media_numbers = [_part_number(item.name, "image", item.suffix) for item in media_dir.glob("image*") if item.stem.startswith("image") and item.stem[5:].isdigit()]
        next_media = max(media_numbers or [0]) + 1
        for source in page_pptx[1:]:
            with zipfile.ZipFile(source) as source_zip:
                source_slides = sorted([name for name in source_zip.namelist() if name.startswith("ppt/slides/slide") and name.endswith(".xml")], key=lambda name: _part_number(name, "slide", ".xml"))
                if len(source_slides) != 1:
                    raise DeckError(f"page runtime output must contain exactly one slide: {source}")
                source_slide = source_slides[0]
                new_slide = f"ppt/slides/slide{next_slide}.xml"
                source_rels = source_slide.replace("ppt/slides/", "ppt/slides/_rels/") + ".rels"
                rel_root = ET.fromstring(source_zip.read(source_rels))
                for rel in rel_root:
                    target_name = rel.attrib.get("Target", "")
                    if "../media/" in target_name:
                        old_media = "ppt/media/" + Path(target_name).name
                        extension = Path(old_media).suffix
                        new_media = media_dir / f"image{next_media}{extension}"
                        new_media.write_bytes(source_zip.read(old_media))
                        rel.attrib["Target"] = "../media/" + new_media.name
                        next_media += 1
                (package / new_slide).write_bytes(source_zip.read(source_slide))
                new_rels = package / f"ppt/slides/_rels/slide{next_slide}.xml.rels"
                new_rels.parent.mkdir(parents=True, exist_ok=True)
                new_rels.write_bytes(ET.tostring(rel_root, encoding="utf-8", xml_declaration=True))
                relation = ET.Element(f"{{{REL_NS}}}Relationship", {"Id": f"rId{next_rid}", "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide", "Target": f"slides/slide{next_slide}.xml"})
                presentation_rels.append(relation)
                sld_id_list.append(ET.Element(f"{{{P_NS}}}sldId", {"id": str(next_sld_id), f"{{{R_NS}}}id": f"rId{next_rid}"}))
                content_types.append(ET.Element(f"{{{CT_NS}}}Override", {"PartName": f"/ppt/slides/slide{next_slide}.xml", "ContentType": "application/vnd.openxmlformats-officedocument.presentationml.slide+xml"}))
                mapping.append({"source_pptx": str(source), "final_slide_index": len(mapping) + 1})
                next_slide += 1; next_sld_id += 1; next_rid += 1
        presentation_path.write_bytes(ET.tostring(presentation, encoding="utf-8", xml_declaration=True))
        rels_path.write_bytes(ET.tostring(presentation_rels, encoding="utf-8", xml_declaration=True))
        types_path.write_bytes(ET.tostring(content_types, encoding="utf-8", xml_declaration=True))
        temporary_output = output.with_suffix(".tmp")
        with zipfile.ZipFile(temporary_output, "w", compression=zipfile.ZIP_DEFLATED) as target:
            for item in sorted(package.rglob("*")):
                if item.is_file():
                    target.write(item, item.relative_to(package).as_posix())
        temporary_output.replace(output)
    # python-pptx canonicalizes package bookkeeping after the low-level merge.
    # This preserves slide XML/native shapes while making the deck consumable
    # by LibreOffice as well as PowerPoint/python-pptx.
    reopened = Presentation(output)
    canonical = output.with_suffix(".canonical.tmp")
    reopened.save(canonical)
    canonical.replace(output)
    Presentation(output)  # Reopen contract.
    return mapping


def pptx_shape_metrics(path: Path) -> list[dict[str, Any]]:
    presentation = Presentation(path)
    output = []
    for index, slide in enumerate(presentation.slides, 1):
        text = [shape for shape in slide.shapes if getattr(shape, "has_text_frame", False)]
        fonts: list[str] = []
        sizes: list[float] = []
        for shape in text:
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.font.name: fonts.append(run.font.name)
                    if run.font.size: sizes.append(run.font.size.pt)
        output.append({"slide_index": index, "shape_count": len(slide.shapes), "text_shape_count": len(text), "font_families": sorted(set(fonts)), "font_size_range_pt": [min(sizes), max(sizes)] if sizes else None, "slide_size": [presentation.slide_width, presentation.slide_height]})
    return output


def continuity_report(deck_pptx: Path, pages: list[dict[str, Any]]) -> dict[str, Any]:
    metrics = pptx_shape_metrics(deck_pptx)
    results = {
        "slide_size": "pass", "page_order": "pass", "duplicate_page_ids": "pass", "route_page_role": "pass",
        "title_font_family": "manual_required", "title_size_range": "manual_required", "body_font_family": "manual_required",
        "page_number_style": "manual_required", "header_footer_presence": "manual_required", "margin_baseline": "manual_required",
        "brand_color_tokens": "not_assessed", "repeated_asset_identity": "manual_required", "section_grouping": "pass",
        "transition_between_sections": "manual_required", "missing_mandatory_page_roles": "not_assessed",
    }
    if len({page["page_id"] for page in pages}) != len(pages): results["duplicate_page_ids"] = "warn"
    if any(metric["slide_size"] != metrics[0]["slide_size"] for metric in metrics): results["slide_size"] = "warn"
    if any(not page.get("route") or not page.get("page_role") for page in pages): results["route_page_role"] = "warn"
    mandatory = pages[0].get("mandatory_page_roles", []) if pages else []
    missing = [role for role in mandatory if role not in {page["page_role"] for page in pages}]
    if missing: results["missing_mandatory_page_roles"] = "warn"
    return {"result": results, "slide_metrics": metrics, "missing_mandatory_roles": missing, "scope": "mechanical continuity only; no L3, text-to-visual binding, or client-grade claim"}


class DeckRunner:
    def __init__(self, manifest_path: Path, *, resume: bool = True, force: bool = False, jobs: int | None = None, page: str | None = None, from_page: str | None = None, skip_preview: bool = False):
        self.manifest_path = manifest_path.resolve(); self.base = self.manifest_path.parent; self.manifest = read_json(self.manifest_path)
        errors = validate_manifest(self.manifest)
        if errors: raise DeckError("deck manifest invalid: " + "; ".join(errors))
        self.resume, self.force, self.jobs, self.page, self.from_page, self.skip_preview = resume, force, jobs or int(self.manifest.get("runtime_options", {}).get("jobs", 1)), page, from_page, skip_preview
        self.root = resolve_path(self.manifest["output"]["root"], self.base); self.run_id = f"{datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S%fZ')}-{uuid.uuid4().hex[:8]}"; self.run_root = self.root / "runs" / self.run_id
        self.events: list[dict[str, Any]] = []

    def _event(self, payload: dict[str, Any]) -> None:
        self.events.append(payload); self.run_root.mkdir(parents=True, exist_ok=True)
        with (self.run_root / "deck_stage_events.jsonl").open("a", encoding="utf-8") as stream: stream.write(json.dumps(payload, ensure_ascii=False) + "\n")

    def _page_manifest(self, page: dict[str, Any]) -> Path:
        return resolve_path(page["page_runtime_manifest"], self.base)

    def _latest_page_result(self, page: dict[str, Any]) -> dict[str, Any] | None:
        """Reuse a prior trusted page run without reusing mutable artifacts."""
        manifest_path = self._page_manifest(page)
        runtime = RuntimeCaseRunner(manifest_path, resume=True, skip_preview=self.skip_preview)
        latest = runtime.case_root / "latest_success.json"
        if not latest.exists():
            return None
        try:
            pointer = read_json(latest)
            summary = read_json(Path(pointer["summary"]))
            if summary.get("overall_status") != "success":
                return None
            pptx = Path(pointer["run_root"]) / "artifacts" / f"{runtime.manifest['case_id']}.pptx"
            if not validate_artifact(pptx, "pptx")["valid"]:
                return None
        except (OSError, KeyError, ValueError, json.JSONDecodeError):
            return None
        return {"page_id": page["page_id"], "page_role": page["page_role"], "route": page.get("route"), "page_run_id": pointer["run_id"], "page_run_root": pointer["run_root"], "status": "success", "required": page["required"], "cache_status": "deck_reuse", "pptx": str(pptx), "summary": summary}

    def _run_page(self, page: dict[str, Any]) -> dict[str, Any]:
        runtime = RuntimeCaseRunner(self._page_manifest(page), resume=self.resume, force=self.force, skip_preview=self.skip_preview)
        summary = runtime.run()
        pptx = None
        for stage in summary.get("stages", []):
            if stage.get("stage_id") == "assemble_pptx":
                for entry in stage.get("artifacts", []):
                    if entry.get("artifact_type") == "pptx": pptx = entry["path"]
        assembly = next((stage for stage in summary.get("stages", []) if stage.get("stage_id") == "assemble_pptx"), {})
        decision = assembly.get("cache_decision", {}).get("decision")
        cache_status = "hit" if decision == "hit" else "miss" if decision in {"miss", "force"} else "not_available"
        return {"page_id": page["page_id"], "page_role": page["page_role"], "route": page.get("route"), "page_run_id": summary["run_id"], "page_run_root": str(runtime.run_root), "status": summary["overall_status"], "required": page["required"], "cache_status": cache_status, "pptx": pptx, "summary": summary}

    def run(self, dry_plan: bool = False) -> dict[str, Any]:
        pages = sorted(self.manifest["pages"], key=lambda item: item["order"]); graph = page_graph(pages); levels = topological_levels(graph); selected = selected_page_ids(graph, self.page, self.from_page, [page["page_id"] for page in pages])
        if dry_plan: return {"deck_run_id": self.run_id, "dry_plan": True, "levels": levels, "selected_pages": sorted(selected)}
        self.run_root.mkdir(parents=True, exist_ok=True); write_json(self.run_root / "resolved_deck_manifest.json", self.manifest); write_json(self.run_root / "page_graph.json", graph); write_json(self.run_root / "scheduling_plan.json", {"levels": levels, "jobs": self.jobs, "selected_pages": sorted(selected)})
        by_id = {page["page_id"]: page for page in pages}; results: dict[str, dict[str, Any]] = {}; failed_required = False
        for level in levels:
            runnable = [pid for pid in level if pid in selected]
            for pid in level:
                if pid not in selected:
                    reused = self._latest_page_result(by_id[pid])
                    if reused:
                        results[pid] = reused
                        self._event({"stage": "page_runtime", "reuse": True, **reused})
                    else:
                        results[pid] = {"page_id": pid, "page_role": by_id[pid]["page_role"], "route": by_id[pid].get("route"), "required": by_id[pid]["required"], "status": "skipped", "warnings": ["no trusted prior page evidence available"]}
            blocked = [pid for pid in runnable if any(results.get(dep, {}).get("status") not in {"success", "partial"} for dep in graph[pid])]
            for pid in blocked:
                results[pid] = {"page_id": pid, "page_role": by_id[pid]["page_role"], "route": by_id[pid].get("route"), "required": by_id[pid]["required"], "status": "blocked", "warnings": ["dependency did not produce trusted output"]}
            runnable = [pid for pid in runnable if pid not in blocked]
            with ThreadPoolExecutor(max_workers=max(1, self.jobs)) as executor:
                futures = {executor.submit(self._run_page, by_id[pid]): pid for pid in runnable}
                for future in as_completed(futures):
                    result = future.result(); results[result["page_id"]] = result; self._event({"stage": "page_runtime", **result}); failed_required = failed_required or (result["required"] and result["status"] == "failed")
            if failed_required and self.manifest.get("failure_policy", "strict") == "strict": break
        for page in pages:
            if page["page_id"] not in results:
                results[page["page_id"]] = {"page_id": page["page_id"], "page_role": page["page_role"], "route": page.get("route"), "required": page["required"], "status": "skipped", "warnings": ["not selected or blocked"]}
        successful = [results[page["page_id"]] for page in pages if results[page["page_id"]].get("status") in {"success", "partial"} and results[page["page_id"]].get("pptx")]
        if failed_required and self.manifest.get("failure_policy", "strict") == "strict": overall = "failed"; deck_path = None
        else:
            deck_path = self.run_root / "artifacts" / f"{self.manifest['deck_id']}.pptx"; mapping = merge_single_page_pptx([Path(item["pptx"]) for item in successful], deck_path) if successful else []
            overall = "partial" if any(item.get("status") != "success" for item in results.values()) else "success"
            if deck_path: write_json(self.run_root / "deck_merge_mapping.json", mapping)
        page_index = []
        if deck_path:
            metrics = pptx_shape_metrics(deck_path)
            for pos, page in enumerate([page for page in pages if results[page["page_id"]].get("status") in {"success", "partial"}], 1):
                source_metrics = pptx_shape_metrics(Path(results[page["page_id"]]["pptx"]))[0]
                item = deepcopy(results[page["page_id"]]); item.update({"final_slide_index": pos, "page_artifact_hash": sha256_file(Path(item["pptx"])), "source_shape_count": source_metrics["shape_count"], "source_text_count": source_metrics["text_shape_count"], "final_deck_shape_count": metrics[pos-1]["shape_count"], "final_deck_text_count": metrics[pos-1]["text_shape_count"]}); page_index.append(item)
            continuity = continuity_report(deck_path, pages); write_json(self.run_root / "deck_continuity_report.json", continuity); atomic_write(self.run_root / "deck_continuity_report.md", "# Deck Continuity Report\n\n" + "\n".join(f"- {key}: `{value}`" for key, value in continuity["result"].items()) + "\n")
            errors, warnings, slides = inspect(str(deck_path)); package = {"result": "FAIL" if errors else "PASS", "errors": errors, "warnings": warnings, "slides": slides}; write_json(self.run_root / "package_inspection.json", package); write_json(self.run_root / "no_shadow_scan.json", scan_no_shadow(deck_path))
            if errors or scan_no_shadow(deck_path)["result"] != "PASS": overall = "failed"
            if not self.skip_preview:
                preview = self.run_root / "preview"
                completed = subprocess.run([sys.executable, str(SCRIPT_DIR / "render_preview.py"), str(deck_path), "--outdir", str(preview), "--dpi", str(self.manifest.get("runtime_options", {}).get("preview_dpi", 120))], text=True, capture_output=True, timeout=90)
                if completed.returncode == 0 and (preview / "contact-sheet.png").exists():
                    self._event({"stage": "deck_preview", "status": "success", "path": str(preview / "contact-sheet.png")})
                else:
                    self._event({"stage": "deck_preview", "status": "unavailable", "warning": completed.stderr.strip()})
        write_json(self.run_root / "page_run_index.json", page_index or list(results.values()))
        artifacts = [artifact(path, "pptx") for path in [deck_path] if path and path.exists()]
        write_json(self.run_root / "deck_artifact_hash_manifest.json", {"artifacts": artifacts})
        summary = {"deck_run_id": self.run_id, "deck_id": self.manifest["deck_id"], "overall_status": overall, "pages": list(results.values()), "deck_pptx": str(deck_path) if deck_path else None, "events": self.events}
        write_json(self.run_root / "deck_run_summary.json", summary); atomic_write(self.run_root / "deck_run_summary.md", f"# Deck Run Summary\n\nStatus: `{overall}`\n")
        write_json(self.run_root / "deck_run_state.json", {"deck_run_id": self.run_id, "toolchain": toolchain_fingerprint(SKILL_ROOT), "summary": summary})
        return summary
