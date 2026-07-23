from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS))

from pptx_primitives import add_primitive  # noqa: E402


def blank_slide():
    prs = Presentation()
    return prs.slides.add_slide(prs.slide_layouts[6])


def test_supported_primitives_have_deterministic_geometry() -> None:
    slide = blank_slide()
    for index, kind in enumerate(("line", "rule", "tick", "dot", "rectangle", "ellipse", "straight_connector", "elbow_connector")):
        _, record = add_primitive(
            slide,
            {
                "id": f"p{index}",
                "type": kind,
                "zone": {"x": 1, "y": 1 + index * 0.1, "w": 1, "h": 0.1},
                "stroke_color": "#123456",
                "fill_color": "#FFFFFF",
                "stroke_width_pt": 1,
            },
        )
        assert record["applied_geometry"]["x"] == 1
        assert record["pptx_shape_name"] == f"p{index}"


def test_unsupported_depth_effect_is_reported() -> None:
    slide = blank_slide()
    _, record = add_primitive(
        slide,
        {"id": "p", "type": "rectangle", "zone": {"x": 0, "y": 0, "w": 1, "h": 1}, "shadow": True},
    )
    assert "shadow" in record["unsupported_fields"]
