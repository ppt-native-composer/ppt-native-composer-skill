#!/usr/bin/env python3
"""Typed, hash-backed validators for incremental runtime artifacts."""

from __future__ import annotations

import hashlib
import json
import zipfile
from pathlib import Path
from typing import Any

from PIL import Image
from pptx import Presentation


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _base(path: Path, artifact_type: str) -> dict[str, Any]:
    return {"path": str(path), "artifact_type": artifact_type, "exists": path.exists(), "valid": False, "sha256": sha256_file(path) if path.exists() and path.is_file() else None, "metadata": {}, "errors": []}


def validate_pptx(path: Path) -> dict[str, Any]:
    result = _base(path, "pptx")
    if not path.exists():
        result["errors"].append("missing")
        return result
    try:
        with zipfile.ZipFile(path) as package:
            names = set(package.namelist())
            for required in ("[Content_Types].xml", "ppt/presentation.xml"):
                if required not in names:
                    result["errors"].append(f"missing required package part: {required}")
            if package.testzip():
                result["errors"].append("corrupt ZIP member")
        presentation = Presentation(path)
        result["metadata"] = {"slide_count": len(presentation.slides), "shape_count": sum(len(slide.shapes) for slide in presentation.slides)}
        if not result["errors"]:
            result["valid"] = True
    except Exception as exc:  # noqa: BLE001
        result["errors"].append(f"pptx validation error: {exc}")
    return result


def validate_json(path: Path, required_keys: list[str] | None = None) -> dict[str, Any]:
    result = _base(path, "json")
    if not path.exists():
        result["errors"].append("missing")
        return result
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
        if not isinstance(payload, dict):
            result["errors"].append("top-level JSON value is not an object")
        else:
            missing = [key for key in required_keys or [] if key not in payload]
            if missing:
                result["errors"].append("missing required keys: " + ", ".join(missing))
            result["metadata"] = {"top_level_keys": sorted(payload), "required_keys": required_keys or []}
        result["valid"] = not result["errors"]
    except Exception as exc:  # noqa: BLE001
        result["errors"].append(f"json validation error: {exc}")
    return result


def validate_image(path: Path) -> dict[str, Any]:
    result = _base(path, "image")
    if not path.exists():
        result["errors"].append("missing")
        return result
    try:
        with Image.open(path) as image:
            image.verify()
        with Image.open(path) as image:
            width, height = image.size
        result["metadata"] = {"width": width, "height": height}
        if width <= 0 or height <= 0:
            result["errors"].append("non-positive dimensions")
        result["valid"] = not result["errors"]
    except Exception as exc:  # noqa: BLE001
        result["errors"].append(f"image validation error: {exc}")
    return result


def validate_text(path: Path) -> dict[str, Any]:
    result = _base(path, "text")
    if not path.exists() or not path.is_file():
        result["errors"].append("missing or not a file")
        return result
    if not path.read_text(encoding="utf-8").strip():
        result["errors"].append("empty text evidence")
    result["valid"] = not result["errors"]
    return result


def validate_artifact(path: Path, artifact_type: str, required_keys: list[str] | None = None) -> dict[str, Any]:
    if artifact_type == "pptx":
        return validate_pptx(path)
    if artifact_type == "json":
        return validate_json(path, required_keys)
    if artifact_type == "image":
        return validate_image(path)
    return validate_text(path)
