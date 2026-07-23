from __future__ import annotations

import sys
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS))

from evc_runtime import RuntimeAudit, render_text_object  # noqa: E402
from ooxml_capabilities import OoxmlCapabilityError, alpha_nodes, alpha_value, set_shape_fill_opacity  # noqa: E402
from pptx_primitives import add_primitive  # noqa: E402


def _context() -> dict:
    return {"fonts": {"cjk": "PingFang SC", "latin": "Aptos"}, "default_color": "#112233", "default_size_pt": 18}


def test_alpha_value_validates_the_closed_unit_interval() -> None:
    assert alpha_value(0.4) == 40000
    with pytest.raises(OoxmlCapabilityError):
        alpha_value(1.01)
    with pytest.raises(OoxmlCapabilityError):
        alpha_value("not-a-number")


def test_native_primitive_alpha_is_singleton_and_round_trips(tmp_path: Path) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rectangle, _ = add_primitive(slide, {"id": "rect", "type": "rectangle", "zone": {"x": 1, "y": 1, "w": 2, "h": 1}, "fill_color": "#224466", "opacity": 0.4})
    set_shape_fill_opacity(rectangle, 0.4)
    connector, _ = add_primitive(slide, {"id": "line", "type": "straight_connector", "geometry": {"type": "axis", "x1": 1, "y1": 3, "x2": 4, "y2": 3}, "stroke_color": "#112233", "opacity": 0.6})
    output = tmp_path / "alpha.pptx"
    prs.save(output)
    Presentation(output)
    with zipfile.ZipFile(output) as package:
        xml = package.read("ppt/slides/slide1.xml")
    values = alpha_nodes(xml)
    assert values.count(40000) == 1
    assert values.count(60000) == 1
    assert connector.name == "line"


def test_text_font_and_textbox_alpha_round_trip(tmp_path: Path) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    item = {
        "id": "text", "source_id": "text", "role": "claim", "text": "透明文字", "zone": {"x": 1, "y": 1, "w": 3, "h": 1},
        "appearance": {"fill": "#FFFFFF", "stroke": "#224466", "opacity": 0.5},
        "typographic_treatment": {"cjk_size_pt": 24, "latin_size_pt": 24, "weight_token": "bold", "color": "#112233", "font_opacity": 0.7, "overflow_policy": {"mode": "warn", "minimum_size_pt": 20}},
    }
    record = render_text_object(slide, item, _context(), RuntimeAudit(1))
    assert record["unsupported_fields"] == []
    output = tmp_path / "text_alpha.pptx"
    prs.save(output)
    Presentation(output)
    with zipfile.ZipFile(output) as package:
        values = alpha_nodes(package.read("ppt/slides/slide1.xml"))
    assert 50000 in values
    assert 70000 in values
