from __future__ import annotations

import json
import shutil
import sys
from pathlib import Path

import pytest
from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

import run_example
from assemble_pptx import assemble


def test_example_is_self_contained_and_reopens(tmp_path: Path) -> None:
    output = tmp_path / "example"
    report = run_example.run_example(output)
    assert report["status"] == "mechanical_pass"
    assert report["editable_text_count"] == 4
    assert report["text_preserved"]
    assert report["visual_qa"] == "not_run"
    assert report["client_grade"] is False
    assert report["package_inspection"]["result"] == "PASS"
    presentation = Presentation(output / "example.pptx")
    assert len(presentation.slides) == 1
    pictures = [shape for shape in presentation.slides[0].shapes if shape.shape_type == 13]
    assert len(pictures) == 1
    assert pictures[0].image.blob == (output / "assets" / "swatch.png").read_bytes()
    assert abs(presentation.slide_width / presentation.slide_height - 16 / 9) < 0.0001


@pytest.mark.parametrize("kind", ["empty_directory", "directory_with_evidence", "file"])
def test_example_refuses_existing_output(tmp_path: Path, kind: str) -> None:
    output = tmp_path / "existing"
    if kind == "file":
        output.write_text("keep", encoding="utf-8")
        sentinel = output
    else:
        output.mkdir()
        sentinel = output / "evidence.json"
        if kind == "directory_with_evidence":
            sentinel.write_text("keep", encoding="utf-8")
    with pytest.raises(FileExistsError):
        run_example.run_example(output)
    if sentinel.exists():
        assert sentinel.read_text(encoding="utf-8") == "keep"
    assert not (output / "example.pptx").exists()


def test_example_does_not_report_success_when_source_copy_disagrees(tmp_path: Path, monkeypatch) -> None:
    fixture = tmp_path / "fixture"
    shutil.copytree(run_example.FIXTURE_DIR, fixture)
    (fixture / "copy.json").write_text('{"title": "different source"}', encoding="utf-8")
    monkeypatch.setattr(run_example, "FIXTURE_DIR", fixture)
    output = tmp_path / "result"
    with pytest.raises(RuntimeError, match="verification failed"):
        run_example.run_example(output)
    assert not (output / "example_result.json").exists()


def test_fixture_does_not_bypass_approval_gate(tmp_path: Path, monkeypatch) -> None:
    fixture = tmp_path / "fixture"
    shutil.copytree(run_example.FIXTURE_DIR, fixture)
    blueprint = json.loads((fixture / "blueprint.json").read_text(encoding="utf-8"))
    blueprint["slides"][0]["editable_visual_composition"]["approval"]["approved"] = False
    (fixture / "blueprint.json").write_text(json.dumps(blueprint), encoding="utf-8")
    monkeypatch.setattr(run_example, "FIXTURE_DIR", fixture)
    output = tmp_path / "result"
    with pytest.raises(RuntimeError, match="approval.approved"):
        run_example.run_example(output)
    assert not (output / "example.pptx").exists()


def test_label_follows_reference_geometry_in_actual_pptx(tmp_path: Path) -> None:
    original = tmp_path / "original"
    run_example.run_example(original)
    baseline_bytes = (original / "example.pptx").read_bytes()
    revised = tmp_path / "revised"
    shutil.copytree(original, revised)
    path = revised / "blueprint.json"
    data = json.loads(path.read_text(encoding="utf-8"))
    slide = data["slides"][0]
    slide["sourced_asset_layer"][0]["expected_zone"]["x"] += 1
    slide["editable_visual_composition"]["visual_targets"][0]["geometry"]["x"] += 1
    # Only the artwork and its declared target change, not the label zone.
    path.write_text(json.dumps(data), encoding="utf-8")
    assemble(path, revised / "moved.pptx", revised)

    def label_left(pptx: Path) -> int:
        return next(shape.left for shape in Presentation(pptx).slides[0].shapes if shape.has_text_frame and shape.text == "Reference swatch")

    # DrawingML positions are integer EMUs; float-to-EMU rounding can differ by one.
    assert label_left(revised / "moved.pptx") - label_left(original / "example.pptx") == pytest.approx(914400, abs=1)
    assert (original / "example.pptx").read_bytes() == baseline_bytes


def test_editable_copy_can_change_without_replacing_artwork(tmp_path: Path) -> None:
    output = tmp_path / "example"
    run_example.run_example(output)
    data = json.loads((output / "blueprint.json").read_text(encoding="utf-8"))
    slide = data["slides"][0]
    text = "A revised editable label"
    slide["editable_layer"]["body"][0]["text"] = text
    next(item for item in slide["editable_visual_composition"]["editable_objects"] if item["id"] == "example_label")["text"] = text
    revised = tmp_path / "revision"
    revised.mkdir()
    shutil.copytree(output / "assets", revised / "assets")
    blueprint = revised / "blueprint.json"
    blueprint.write_text(json.dumps(data), encoding="utf-8")
    assemble(blueprint, revised / "example.pptx", revised)
    presentation = Presentation(revised / "example.pptx")
    assert text in [shape.text for shape in presentation.slides[0].shapes if shape.has_text_frame]
    assert next(shape.image.blob for shape in presentation.slides[0].shapes if shape.shape_type == 13) == (output / "assets" / "swatch.png").read_bytes()
